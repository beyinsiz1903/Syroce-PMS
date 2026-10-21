"""Bounded file readers. No remote URLs, database dumps, macros or formulas."""

from __future__ import annotations

import csv
import io
import json
import math
import zipfile
from datetime import date, datetime
from pathlib import Path

from defusedxml import ElementTree

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 80
MAX_CELL = 4096


class MigrationError(ValueError):
    pass


def column_key(value):
    import re
    import unicodedata

    value = str(value).translate(str.maketrans({"ı": "i", "İ": "I"}))
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", value)


FORBIDDEN_COLUMNS = {
    "password",
    "sifre",
    "clientsecret",
    "accesstoken",
    "refreshtoken",
    "cvv",
    "cvc",
    "creditcard",
    "creditcardnumber",
    "cardnumber",
    "kartnumarasi",
    "kredikarti",
    "pan",
}


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (dict, list, tuple)) or isinstance(value, float) and not math.isfinite(value):
        raise MigrationError("İç içe veya sonlu olmayan hücre değeri desteklenmiyor.")
    value = str(value).strip()
    if len(value) > MAX_CELL or "\x00" in value:
        raise MigrationError("Hücre çok uzun veya geçersiz karakter içeriyor.")
    return value


def _table(headers, rows):
    headers = [_cell(h) for h in headers]
    if not headers or len(headers) > MAX_COLUMNS or any(not h or len(h) > 160 for h in headers):
        raise MigrationError("1–80 adet, boş olmayan sütun başlığı gerekli.")
    normalized = [column_key(h) for h in headers]
    if any(not h for h in normalized) or len(set(normalized)) != len(headers):
        raise MigrationError("Tekrarlanan veya geçersiz sütun başlığı.")
    if set(normalized) & FORBIDDEN_COLUMNS:
        raise MigrationError("Şifre, token veya ödeme kartı sütunlarını dosyadan kaldırın.")
    output = []
    for row in rows:
        if len(row) != len(headers):
            raise MigrationError("Satırdaki hücre sayısı başlıklarla uyuşmuyor.")
        values = [_cell(v) for v in row]
        if not any(values):
            continue
        output.append(dict(zip(headers, values, strict=True)))
        if len(output) > MAX_ROWS:
            raise MigrationError("Dosya başına en fazla 10.000 satır; dosyayı parçalara ayırın.")
    if not output:
        raise MigrationError("Dosyada veri satırı yok.")
    return headers, output


def parse_file(filename, content, *, encoding="utf-8-sig", sheet=None):
    if not content or len(content) > MAX_BYTES:
        raise MigrationError("Dosya boş veya 5 MB sınırını aşıyor.")
    extension = Path(filename or "").suffix.lower()
    try:
        if extension == ".xlsx":
            from openpyxl import load_workbook

            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(e.file_size for e in entries) > 40 * 1024 * 1024:
                    raise MigrationError("Excel açılmış boyut sınırını aşıyor.")
                if any("vbaProject" in e.filename or "externalLinks/" in e.filename for e in entries):
                    raise MigrationError("Makro veya dış bağlantı içeren dosya kabul edilmez.")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if sheet is None and len(book.sheetnames) != 1:
                    raise MigrationError("Çok sayfalı Excel için sayfa adını belirtin.")
                if sheet is not None and sheet not in book.sheetnames:
                    raise MigrationError("Excel sayfası bulunamadı.")
                page = book[sheet or book.sheetnames[0]]
                if page.max_row > MAX_ROWS + 1 or page.max_column > MAX_COLUMNS:
                    raise MigrationError("Excel satır veya sütun sınırını aşıyor.")

                def values():
                    for row in page.iter_rows():
                        if any(cell.data_type == "f" for cell in row):
                            raise MigrationError("Excel formüllerini önce sabit değerlere dönüştürün.")
                        yield [cell.value for cell in row]

                iterator = values()
                return _table(next(iterator, []), iterator)
            finally:
                book.close()
        if encoding not in {"utf-8-sig", "cp1254"}:
            raise MigrationError("Desteklenmeyen metin kodlaması.")
        text = content.decode(encoding, errors="strict")
        if extension in {".csv", ".tsv"}:
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel_tab if extension == ".tsv" else csv.excel
            reader = csv.reader(io.StringIO(text), dialect=dialect, strict=True)
            return _table(next(reader, []), reader)
        if extension == ".json":

            def unique_object(pairs):
                result = {}
                for key, value in pairs:
                    if key in result:
                        raise MigrationError("JSON içinde tekrarlanan alan adı.")
                    result[key] = value
                return result

            records = json.loads(text, object_pairs_hook=unique_object)
            if not isinstance(records, list) or not records or not all(isinstance(r, dict) for r in records):
                raise MigrationError("JSON, düz nesnelerden oluşan bir liste olmalı.")
            headers = list(dict.fromkeys(k for row in records for k in row))
            return _table(headers, ([row.get(h) for h in headers] for row in records))
        if extension == ".xml":
            root = ElementTree.fromstring(text)
            records = list(root)
            if not records or len(records) > MAX_ROWS:
                raise MigrationError("XML kökünde 1–10.000 kayıt gerekli.")
            headers = list(dict.fromkeys(child.tag for row in records for child in row))
            rows = []
            for row in records:
                if row.attrib or any(list(child) or child.attrib for child in row):
                    raise MigrationError("XML düz alanlardan oluşmalı; iç içe/öznitelikli yapı için önce dönüştürün.")
                if len({child.tag for child in row}) != len(row):
                    raise MigrationError("XML kaydında tekrarlanan alan.")
                data = {child.tag: child.text for child in row}
                rows.append([data.get(h) for h in headers])
            return _table(headers, rows)
        raise MigrationError("CSV, TSV, XLSX, JSON veya düz XML dosyası seçin.")
    except MigrationError:
        raise
    except Exception as exc:
        # Parser exceptions may quote a cell containing PII; never expose them.
        raise MigrationError("Dosya okunamadı. Biçimi, kodlamayı ve hücreleri kontrol edin.") from exc
