"""Presentation only: keep the occupancy calculator and snapshot contract intact."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


def occupancy_summary(data):
    """One clearly labelled metric per date, without counting today twice."""
    rows = []
    for row in data["rows"]:
        actual = row["date"] < data["business_date"]
        prefix = "actual" if actual else "expected"
        occupied = row[f"{prefix}_rooms"]
        rows.append({
            "date": row["date"], "occupancy": row[f"{prefix}_occupancy"],
            "occupied": occupied, "available": row["available_rooms"],
            "vacant": max(0, row["available_rooms"] - occupied) if occupied is not None else None,
            "basis": "Gerçekleşen" if actual else "Beklenen",
        })
    known = [row for row in rows if row["occupied"] is not None]
    nights = sum(row["occupied"] for row in known)
    capacity = sum(row["available"] for row in known)
    return rows, {"room_nights": nights, "available_room_nights": capacity,
                  "occupancy": nights / capacity if capacity else None}


def occupancy_workbook(data):
    """Compact operator view first; complete reconciliation data stays inspectable."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Doluluk Özeti"
    rows, total = occupancy_summary(data)
    first, last = (date.fromisoformat(data[key]).strftime("%d.%m.%Y") for key in ("start_date", "end_date"))
    scope = data["room_type"] or "Tüm tesis"
    summary.append(["Günlük Tesis Doluluğu"])
    summary.append([f"{first} – {last}   /   {scope}"])
    summary.append([f"PMS iş günü: {date.fromisoformat(data['business_date']).strftime('%d.%m.%Y')}. Öncesi gerçekleşen; iş günü ve sonrası beklenen doluluk."])
    summary.append(["Dönem doluluğu", total["occupancy"], "Dolu oda-gecesi", total["room_nights"], "Satılabilir oda-gecesi", total["available_room_nights"]])
    summary.append([])
    summary.append(["Tarih", "Doluluk %", "Dolu oda", "Boş oda", "Satılabilir oda", "Veri türü"])
    for row in rows:
        summary.append([date.fromisoformat(row["date"]), row["occupancy"] / 100 if row["occupancy"] is not None else None,
                        row["occupied"], row["vacant"], row["available"], row["basis"]])
    if any(row.get(key) for row in data["rows"] for key in
           ("assignment_conflicts", "unknown_room_references", "occupied_blocked_rooms", "overbooked_rooms")):
        summary.append([])
        summary.append(["Uyarı: envanter tutarsızlığı veya fazla satış var. Tarih ve nedenler Detaylar sayfasındadır; oranlar %100 ile sınırlandırılmadı."])
        summary.merge_cells(start_row=summary.max_row, start_column=1, end_row=summary.max_row, end_column=6)
        summary.row_dimensions[summary.max_row].height = 32
    for index in (1, 2, 3):
        summary.merge_cells(start_row=index, start_column=1, end_row=index, end_column=6)
    summary.row_dimensions[3].height = 30
    summary.cell(3, 1).alignment = Alignment(wrap_text=True, vertical="center")
    summary.cell(4, 2).number_format = "0.00%"
    for row in summary.iter_rows(min_row=7, max_row=6 + len(rows)):
        row[0].number_format = "dd.mm.yyyy"
        row[1].number_format = "0.00%"
    summary.cell(4, 2).font = Font(name="Arial", size=12, bold=True, color="1D4ED8")
    summary.conditional_formatting.add(f"B7:B{6 + len(rows)}", CellIsRule(operator="greaterThan", formula=[1],
                                       fill=PatternFill("solid", fgColor="FEE2E2")))

    detail = workbook.create_sheet("Detaylar")
    detail.append(["Doluluk Hesaplama Detayları"])
    detail.append([f"{first} – {last}   /   {scope}   /   PMS iş günü: {data['business_date']}"])
    detail.append(["Beklenen = devam eden + giriş yapacak kesin rezervasyonlar. Çıkış gecesi, iptal/no-show ve opsiyonlar hariç."])
    detail.append(["Kapasite güncel aktif oda envanterine dayanır; tarihsel envanter değildir. Atanmamış kesin rezervasyonlar beklenene dahildir."])
    detail.append([])
    fields = ["date", "physical_rooms", "blocked_rooms", "available_rooms", "actual_rooms", "actual_occupancy",
              "stayovers", "arrivals", "expected_rooms", "expected_occupancy", "unassigned_rooms", "overbooked_rooms",
              "assignment_conflicts", "unknown_room_references", "occupied_blocked_rooms"]
    detail.append(["Tarih", "Fiziksel oda", "Satışa kapalı", "Satılabilir oda", "Gerçekleşen oda", "Gerçekleşen %",
                   "Devam eden", "Giriş yapacak / yapan", "Beklenen oda", "Beklenen %", "Atanmamış", "Fazla satış",
                   "Atama çakışması", "Bilinmeyen oda", "Kapalı odada konaklama"])
    for row in data["rows"]:
        detail.append([date.fromisoformat(row[key]) if key == "date" else
                       row[key] / 100 if key.endswith("_occupancy") and row[key] is not None else row[key]
                       for key in fields])
    for index in (1, 2, 3, 4):
        detail.merge_cells(start_row=index, start_column=1, end_row=index, end_column=15)
    for row in detail.iter_rows(min_row=7, max_row=6 + len(rows)):
        row[0].number_format = "dd.mm.yyyy"
        for column in (5, 9):
            row[column].number_format = "0.00%"
    detail.append([])
    for key, label in (("actual", "Gerçekleşen"), ("expected", "Beklenen")):
        item = data["totals"][key]
        detail.append([label, "Dolu oda-gecesi", item["room_nights"], "Satılabilir oda-gecesi", item["available_room_nights"],
                       "Doluluk %", item["occupancy"] / 100 if item["occupancy"] is not None else None])
        detail.cell(detail.max_row, 7).number_format = "0.00%"

    # Preserve text, including formula-looking room-type names, without exposing PII.
    for sheet, widths in ((summary, [23, 16, 23, 15, 28, 21]), (detail, [17] + [24] * 14)):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B7"
        sheet.auto_filter.ref = f"A6:{get_column_letter(sheet.max_column)}{6 + len(rows)}"
        sheet.print_title_rows = "1:6"
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1 if sheet == summary else 0
        sheet.page_setup.fitToHeight = 0
        sheet.print_area = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        sheet.page_margins = PageMargins(left=.25, right=.25, top=.4, bottom=.4, header=.2, footer=.2)
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows():
            sheet.row_dimensions[row[0].row].height = max(sheet.row_dimensions[row[0].row].height or 0, 23)
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
                cell.font = Font(name="Arial", size=10, color="0F172A")
                cell.alignment = Alignment(vertical="center", horizontal="left" if isinstance(cell.value, str) else "right")
        sheet.cell(1, 1).font = Font(name="Arial", size=15, bold=True, color="0F172A")
        for cell in sheet[6]:
            cell.fill = PatternFill("solid", fgColor="1E3A5F")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[6].height = 32
        for index in range(7, 7 + len(rows)):
            if index % 2:
                for cell in sheet[index]:
                    cell.fill = PatternFill("solid", fgColor="F1F5F9")
    summary.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    summary.cell(3, 1).alignment = Alignment(wrap_text=True, vertical="center")
    if summary.max_row > 6 + len(rows):
        summary.cell(summary.max_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        summary.cell(summary.max_row, 1).font = Font(name="Arial", size=10, color="9A3412")
    summary.cell(4, 2).font = Font(name="Arial", size=12, bold=True, color="1D4ED8")
    for column in (2, 4, 6):
        summary.cell(4, column).alignment = Alignment(horizontal="center", vertical="center")
    for index in range(7, 7 + len(rows)):
        summary.cell(index, 6).alignment = Alignment(horizontal="center", vertical="center")
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
