"""Forecast table definitions shared with JSON clients and same-snapshot XLSX."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FIELDS = {
    "physical_rooms": "Oda kapasitesi", "available_rooms": "Satılabilir oda", "vacant_rooms": "Müsait oda",
    "booked_rooms": "Dolu oda", "block_rooms": "Blok oda (kalan kesin)", "total_occupied_rooms": "Toplam dolu oda",
    "blocked_rooms": "Kapalı oda", "remaining_rooms": "Net kalan (+/−)", "occupancy_pct": "Doluluk % (fiziksel)",
    "sellable_occupancy_pct": "Doluluk % (satılabilir)", "arrivals": "Gelen oda", "departures": "Ayrılan oda", "stayovers": "Devam eden",
    "adults": "Yetişkin", "children": "Çocuk / yaşı eksik", "infants": "Bebek (0–2)", "unknown_child_ages": "Yaşı eksik çocuk",
    "comp_rooms": "Comp oda", "day_use_rooms": "Day-use oda", "house_use_rooms": "House-use oda", "group_rooms": "Pickup grup oda",
    "tentative_block_rooms": "Opsiyon blok oda", "unassigned_rooms": "Atanmamış oda", "overbooked_rooms": "Fazla satış",
    "room_revenue_by_currency": "Rezervasyon oda bedeli", "adr_by_currency": "Ort. ücretli oda bedeli",
}
VIEWS = {
    "availability": ["physical_rooms", "vacant_rooms", "total_occupied_rooms", "adults", "children", "infants", "booked_rooms", "block_rooms", "blocked_rooms", "remaining_rooms", "comp_rooms", "day_use_rooms", "house_use_rooms", "group_rooms", "arrivals", "departures", "room_revenue_by_currency", "adr_by_currency", "occupancy_pct"],
    "movements": ["arrivals", "departures", "stayovers", "booked_rooms", "block_rooms", "total_occupied_rooms", "adults", "children", "infants", "unassigned_rooms"],
    "occupancy": ["physical_rooms", "available_rooms", "booked_rooms", "block_rooms", "total_occupied_rooms", "remaining_rooms", "occupancy_pct", "sellable_occupancy_pct", "overbooked_rooms"],
    "revenue": ["booked_rooms", "comp_rooms", "room_revenue_by_currency", "adr_by_currency", "occupancy_pct"],
    "stay_types": ["booked_rooms", "comp_rooms", "house_use_rooms", "day_use_rooms", "group_rooms", "block_rooms", "tentative_block_rooms", "unknown_child_ages"],
}


def forecast_columns():
    return {key: [{"key": field, "label": FIELDS[field], "type": "money" if field.endswith("_by_currency") else "percent" if field.endswith("_pct") else "number"} for field in fields] for key, fields in VIEWS.items()}


def forecast_workbook(data, view, basis="physical"):
    book = Workbook()
    sheet = book.active
    sheet.title = "Forecast"
    sheet.append(["Ön Büro Forecast — Kayıtlı Rezervasyon ve Grup Blokları"])
    sheet.append([f"{data['start_date']} – {data['end_date']} / {data['room_type'] or 'Tüm tesis'} / İş günü: {data['business_date']}"])
    sheet.append([data["definitions"]["occupancy"]])
    sheet.append([data["definitions"]["revenue"]])
    sheet.append([f"Toplam dolu oda-gecesi: {data['totals']['occupied_room_nights'] if data['totals']['occupied_room_nights'] is not None else 'Doğrulanamadı'} / Doluluk bazı: {'Fiziksel oda kapasitesi' if basis == 'physical' else 'Satılabilir oda kapasitesi'}"])
    columns = [("date", "Tarih", None)]
    for key in VIEWS[view]:
        if key == "occupancy_pct" and basis == "sellable" and view != "occupancy":
            key = "sellable_occupancy_pct"
        if key.endswith("_by_currency"):
            currencies = sorted({code for row in data["rows"] for code in (row.get(key) or {})}) or [data["currency"]]
            columns.extend((key, f"{FIELDS[key]} ({code})", code) for code in currencies)
        else:
            columns.append((key, FIELDS[key], None))
    sheet.append([label for _, label, _ in columns])
    for row in data["rows"]:
        sheet.append([date.fromisoformat(row[key]) if key == "date" else
                      (row[key].get(code, 0) if row.get(key) is not None else None) if code else
                      row.get(key) / 100 if key.endswith("_pct") and row.get(key) is not None else row.get(key)
                      for key, _, code in columns])
    last_column = len(columns)
    for index in range(1, 6):
        sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=last_column)
        sheet.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[index].height = 28 if index != 4 else 42
    for cell in sheet[6]:
        cell.fill = PatternFill("solid", fgColor="123047")
        cell.font = Font(color="FFFFFF", bold=True)
    for index, (key, _, code) in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 14 if key == "date" else 23 if code else 18
        for row in sheet.iter_rows(min_row=7, min_col=index, max_col=index):
            cell = row[0]
            cell.number_format = "dd.mm.yyyy" if key == "date" else "0.00%" if key.endswith("_pct") else "#,##0.00" if code else "+0;-0;0" if key == "remaining_rooms" else "0"
            if key == "remaining_rooms" and cell.value is not None and cell.value < 0:
                cell.fill = PatternFill("solid", fgColor="FEE2E2")
    notes = book.create_sheet("Tanımlar ve Uyarılar")
    notes.append(["Hesaplama tanımları"])
    for value in data["definitions"].values():
        notes.append([value])
    for warning in data["warnings"]:
        notes.append([warning])
    for row in data["rows"]:
        if any(row.get(key) for key in ("unallocated_group_blocks", "assignment_conflicts", "unknown_room_references", "occupied_blocked_rooms", "overbooked_rooms")):
            notes.append([f"{row['date']}: {row['unallocated_group_blocks']} belirsiz blok, {row['assignment_conflicts']} atama çakışması, {row['unknown_room_references']} bilinmeyen oda, {row['occupied_blocked_rooms']} kapalı odada konaklama, {row['overbooked_rooms']} fazla satış."])
    notes.column_dimensions['A'].width = 130
    for row in notes:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        notes.row_dimensions[row[0].row].height = 45
    graph = book.create_sheet("Müsaitlik Grafiği")
    graph.append(["Tarih", "Rezervasyon", "Kesin grup bloğu (kalan)", "Müsait", "Kapalı"])
    for row in data["rows"]:
        graph.append([date.fromisoformat(row['date']).strftime('%d.%m.%Y'), row['booked_rooms'], row['block_rooms'], row['vacant_rooms'], row['blocked_rooms']])
    if data["rows"]:
        chart = BarChart()
        chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
        chart.title, chart.y_axis.title = "Günlük Müsaitlik", "Oda"
        chart.width, chart.height = 32, 13
        chart.add_data(Reference(graph, min_col=2, max_col=5, min_row=1, max_row=len(data["rows"]) + 1), titles_from_data=True)
        chart.set_categories(Reference(graph, min_col=1, min_row=2, max_row=len(data["rows"]) + 1))
        for series, color in zip(chart.series, ["38BDF8", "818CF8", "FDE68A", "64748B"], strict=True):
            series.graphicalProperties.solidFill = color
        graph.add_chart(chart, "G2")
    for page in book:
        page.sheet_view.showGridLines = False
        for row in page:
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    sheet.freeze_panes = "B7"
    sheet.auto_filter.ref = f"A6:{get_column_letter(last_column)}{sheet.max_row}"
    sheet.print_title_rows = "1:6"
    sheet.page_setup.orientation, sheet.page_setup.paperSize = "landscape", sheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer
