"""Auto-split from reports.py — backward-compatible sub-router."""

import asyncio
import logging
import uuid
from datetime import UTC, datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer

security = HTTPBearer()

from core.business_date_service import accounting_day_match, accounting_period_match, ensure_business_date_initialized
from core.database import db
from core.email import send_email
from core.helpers import require_module
from core.security import get_current_user
from models.schemas import User
from modules.pms_core.role_permission_service import require_op
from modules.pms_core.stay_night_metrics import NON_COMMERCIAL_STATUSES, as_date, booking_occupies_night, load_stay_night_metrics

try:
    from domains.pms.night_audit_module import AuditStatus, AutomaticPosting, NightAuditRecord
except ImportError:
    NightAuditRecord = None
    AuditStatus = None
    AutomaticPosting = None

from core.utils import (
    create_excel_workbook,
    excel_response,
)
from shared_kernel.migration_observability import migration_observability_service

try:
    from infra.logging_service import get_logging_service
except ImportError:
    get_logging_service = None

try:
    from cache_manager import cached
except ImportError:

    def cached(ttl=300, key_prefix=""):
        def decorator(func):
            return func

        return decorator


logger = logging.getLogger(__name__)
sub_router = APIRouter()

FNB_CATEGORIES = {
    "alcohol",
    "alcoholic_beverage",
    "appetizer",
    "bar",
    "beverage",
    "cafe",
    "dessert",
    "drink",
    "fb",
    "f&b",
    "fnb",
    "food",
    "food_and_beverage",
    "food_beverage",
    "restaurant",
    "room_service",
}


def _report_date(value: str | None, field: str = "date"):
    if not value:
        return None
    try:
        return datetime.fromisoformat(value[:10]).date()
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=422, detail=f"{field} YYYY-MM-DD formatında olmalı") from exc


@sub_router.get("/reports/migration-observability")
async def get_migration_observability(
    current_user: User = Depends(get_current_user),
    _: None = Depends(require_module("reports")),
):
    return await migration_observability_service.get_dashboard(current_user.tenant_id)


@sub_router.post("/reports/send-flash-now")
async def send_flash_report_now(
    recipients: list[str],
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_reports")),  # v98 DW
):
    """Flash report'u şimdi gönder"""
    from modules.analytics_export.report_automation import get_report_automation
    from modules.messaging.email_service import email_service

    automation = get_report_automation(db, email_service)
    await automation.send_flash_report_email(current_user.tenant_id, recipients)

    return {"success": True, "message": f"Flash report {len(recipients)} alıcıya gönderildi"}


@sub_router.get("/reports/flash-report")
@cached(ttl=15, key_prefix="flash_report")  # Cache for 15s
async def get_flash_report(
    date: str | None = None,
    current_user: User = Depends(get_current_user),
    _: None = Depends(require_module("reports")),
    _perm=Depends(require_op("view_reports")),  # v71 Bug DH
):
    """
    Daily Flash Report - Günlük özet rapor
    5 yıldızlı otel yöneticileri için sabah raporu
    """
    if date:
        target_day = _report_date(date)
    else:
        business_state = await ensure_business_date_initialized(db, current_user.tenant_id)
        target_day = datetime.fromisoformat(business_state["business_date"][:10]).date()
    target_key = target_day.isoformat()
    next_key = (target_day + timedelta(days=1)).isoformat()

    metrics = await load_stay_night_metrics(db, current_user.tenant_id, target_day, target_day, actual_only=True)
    metric = metrics[0] if metrics else {"occupied_rooms": 0, "total_rooms": 0, "occupancy_rate": 0}
    total_rooms = metric["total_rooms"]
    occupied_today = metric["occupied_rooms"]
    occupancy_rate = metric["occupancy_rate"]

    day_bookings = await db.bookings.find(
        {
            "tenant_id": current_user.tenant_id,
            "$or": [
                {"check_in": {"$gte": target_key, "$lt": next_key}},
                {"check_out": {"$gte": target_key, "$lt": next_key}},
                {"check_in": {"$lte": target_key}, "check_out": {"$gt": target_key}},
            ],
        },
        {"_id": 0},
    ).to_list(10000)
    arrivals_today = sum(1 for booking in day_bookings if as_date(booking.get("check_in")) == target_day and str(booking.get("status") or "").lower() not in NON_COMMERCIAL_STATUSES)
    departures_today = sum(1 for booking in day_bookings if as_date(booking.get("check_out")) == target_day and str(booking.get("status") or "").lower() not in NON_COMMERCIAL_STATUSES)
    in_house_bookings = [booking for booking in day_bookings if booking_occupies_night(booking, target_day, actual_only=True)]
    inhouse_count = occupied_today

    room_ids = list({b.get("room_id") for b in in_house_bookings if b.get("room_id")})
    room_map = {}
    if room_ids:
        rooms = await db.rooms.find(
            {"tenant_id": current_user.tenant_id, "id": {"$in": room_ids}},
            {"_id": 0, "id": 1, "room_number": 1, "room_no": 1, "name": 1}
        ).to_list(None)
        for r in rooms:
            room_map[r.get("id")] = r.get("room_number") or r.get("room_no") or r.get("name") or "?"

    total_revenue = 0
    collected = 0
    charges_by_cat = {}
    room_revenue_breakdown = []

    for b in in_house_bookings:
        try:
            ci = as_date(b.get("check_in"))
            co = as_date(b.get("check_out"))
            nights = max(1, (co - ci).days)
        except Exception:
            nights = 1

        daily_amount = float(b.get("total_amount") or 0) / nights
        daily_paid = float(b.get("paid_amount") or 0) / nights

        total_revenue += daily_amount
        collected += daily_paid

        room_revenue_breakdown.append({
            "guest_name": b.get("guest_name", "Misafir"),
            "room_number": str(b.get("room_number") or room_map.get(b.get("room_id")) or "?").strip() or "?",
            "daily_rate": round(daily_amount, 2),
            "total_stay_amount": b.get("total_amount", 0),
            "nights": nights
        })


        for c in b.get("charges", []):
            cat = c.get("charge_category", "other")
            amt = float(c.get("amount") or 0) / nights
            charges_by_cat[cat] = charges_by_cat.get(cat, 0) + amt

    # Accounting revenue is authoritative from posted, non-void folio charges.
    # Reservation totals above are kept only for the per-room explanatory list.
    posted_charges = await db.folio_charges.find(
        {
            "tenant_id": current_user.tenant_id,
            "voided": {"$ne": True},
            "$or": [
                {"business_date": target_key},
                {"business_date": {"$exists": False}, "date": {"$regex": f"^{target_key}"}},
                {"business_date": None, "date": {"$regex": f"^{target_key}"}},
            ],
        },
        {"_id": 0},
    ).to_list(10000)
    if posted_charges:
        charges_by_cat = {}
        for charge in posted_charges:
            category = str(charge.get("charge_category") or charge.get("charge_type") or "other").lower()
            charges_by_cat[category] = charges_by_cat.get(category, 0.0) + float(charge.get("total") or charge.get("amount") or 0)
        total_revenue = sum(charges_by_cat.values())

    payments = await db.payments.find(
        {
            "tenant_id": current_user.tenant_id,
            **accounting_day_match(
                target_key,
                {"payment_date": target_key},
                {"date": target_key},
                {"processed_at": {"$regex": f"^{target_key}"}},
                {"created_at": {"$regex": f"^{target_key}"}},
            ),
        },
        {"_id": 0},
    ).to_list(10000)
    collected = 0.0
    for payment in payments:
        status = str(payment.get("status") or "paid").lower()
        if payment.get("voided") or status in {"void", "voided", "failed", "cancelled", "rejected"}:
            continue
        amount = float(payment.get("amount") or 0)
        if str(payment.get("payment_type") or "").lower() == "refund" and amount > 0:
            amount = -amount
        collected += amount

    no_show_bookings = [
        booking for booking in day_bookings
        if as_date(booking.get("check_in")) == target_day
        and str(booking.get("status") or "").lower() in {"no_show", "noshow"}
    ]
    no_shows = len(no_show_bookings)
    target_start = datetime.combine(target_day, datetime.min.time(), tzinfo=UTC)
    target_end = target_start + timedelta(days=1)
    cancellation_filter = {
            "tenant_id": current_user.tenant_id,
            "status": {"$in": ["cancelled", "canceled"]},
            "$or": [
                {"cancelled_at": {"$regex": f"^{target_key}"}},
                {"cancelled_at": {"$gte": target_start, "$lt": target_end}},
                {"cancelled_at": {"$exists": False}, "updated_at": {"$regex": f"^{target_key}"}},
                {"cancelled_at": None, "updated_at": {"$regex": f"^{target_key}"}},
                {"cancelled_at": {"$exists": False}, "updated_at": {"$gte": target_start, "$lt": target_end}},
                {"cancelled_at": None, "updated_at": {"$gte": target_start, "$lt": target_end}},
            ],
        }
    cancelled_bookings = await db.bookings.find(
        cancellation_filter,
        {
            "_id": 0,
            "id": 1,
            "booking_number": 1,
            "guest_name": 1,
            "room_number": 1,
            "check_in": 1,
            "check_out": 1,
            "cancelled_at": 1,
            "cancellation_reason": 1,
            "channel": 1,
        },
    ).to_list(1000)
    cancellations = len(cancelled_bookings)

    walk_ins = sum(1 for b in day_bookings if as_date(b.get("check_in")) == target_day and str(b.get("channel") or b.get("booking_source") or "").lower() == "walk_in")
    overstays = 0

    # POS revenue follows the hotel accounting day. Query failures must remain
    # visible instead of silently producing a plausible but false zero.
    fnb_orders = await db.pos_orders.find(
        {
            "tenant_id": current_user.tenant_id,
            "status": {"$nin": ["cancelled", "canceled", "void", "voided", "refunded"]},
            **accounting_day_match(
                target_key,
                {"closed_at": {"$regex": f"^{target_key}"}},
                {"created_at": {"$regex": f"^{target_key}"}},
            ),
        },
        {"_id": 0, "total_amount": 1, "grand_total": 1},
    ).to_list(5000)
    fnb_revenue = sum(float(order.get("grand_total") or order.get("total_amount") or 0) for order in fnb_orders)

    room_revenue = charges_by_cat.get("room", charges_by_cat.get("accommodation", charges_by_cat.get("room_charge", 0)))
    if not charges_by_cat:
        room_revenue = total_revenue
    posted_fb_revenue = sum(charges_by_cat.get(category, 0) for category in FNB_CATEGORIES)
    if posted_fb_revenue:
        fnb_revenue = posted_fb_revenue
    spa_revenue = charges_by_cat.get("spa", 0)
    minibar_revenue = charges_by_cat.get("minibar", 0)
    laundry_revenue = charges_by_cat.get("laundry", 0)

    grand_total = total_revenue if posted_charges else total_revenue + fnb_revenue
    other_revenue = max(0, grand_total - room_revenue - fnb_revenue - spa_revenue - minibar_revenue - laundry_revenue)
    # ADR and RevPAR are room-revenue metrics; ancillary revenue must not
    # inflate them.
    adr = room_revenue / occupied_today if occupied_today > 0 else 0
    revpar = room_revenue / total_rooms if total_rooms > 0 else 0

    return {
        "date": target_key,
        "occupancy": {
            "rate": round(occupancy_rate, 2),
            "occupied": occupied_today,
            "total": total_rooms,
            "available": total_rooms - occupied_today,
        },
        "kpi": {
            "adr": round(adr, 2),
            "revpar": round(revpar, 2),
        },
        "revenue": {
            "total": round(grand_total, 2),
            "room": round(room_revenue, 2),
            "fb": round(fnb_revenue, 2),
            "spa": round(spa_revenue, 2),
            "minibar": round(minibar_revenue, 2),
            "laundry": round(laundry_revenue, 2),
            "other": round(other_revenue, 2),
            "collected": round(collected, 2),
            "outstanding": round(grand_total - collected, 2),
            "room_revenue_breakdown": room_revenue_breakdown,
            "basis": "posted_folio_charges" if posted_charges else "allocated_reservation_value",
        },
        "operations": {
            "arrivals": arrivals_today,
            "departures": departures_today,
            "inhouse": inhouse_count,
            "no_shows": no_shows,
            "walk_ins": walk_ins,
            "cancellations": cancellations,
            "overstays": overstays,
        },
        "attention_details": {
            "cancellations": cancelled_bookings,
            "no_shows": [
                {
                    key: booking.get(key)
                    for key in (
                        "id", "booking_number", "guest_name", "room_number",
                        "check_in", "check_out", "channel",
                    )
                }
                for booking in no_show_bookings
            ],
        },
        "scope": {
            "business_date": target_key,
            "occupancy": "Seçili iş gecesinde dolu ve satılabilir odalar",
            "revenue": "Seçili iş gününde kaydedilen folyo gelirleri",
            "collections": "Seçili iş gününde alınan ödemeler; başka konaklamalara ait avansları içerebilir",
            "operations": "Seçili iş günündeki giriş, çıkış ve durum hareketleri",
        },
        "departments": [
            {"name": "Oda Geliri", "amount": round(room_revenue, 2)},
            {"name": "Yiyecek & İçecek", "amount": round(fnb_revenue, 2)},
            {"name": "Spa & Wellness", "amount": round(spa_revenue, 2)},
            {"name": "Minibar", "amount": round(minibar_revenue, 2)},
            {"name": "Çamaşırhane", "amount": round(laundry_revenue, 2)},
            {"name": "Diğer", "amount": round(other_revenue, 2)},
        ],
    }


@sub_router.get("/reports/daily-flash-pdf")
@cached(ttl=600, key_prefix="report_daily_flash_pdf")  # Cache for 10 min
async def get_daily_flash_pdf(
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_reports")),  # v85 DU: daily flash GM/CFO
):
    """
    Export daily flash report as PDF
    """
    from fastapi.responses import Response

    try:
        flash_data = await get_daily_flash_report(None, current_user)
        from core.tenant_currency import get_tenant_currency

        _, currency_symbol = await get_tenant_currency(current_user.tenant_id)

        html_content = f"""
        <html>
        <head>
            <style>
                body {{ font-family: Arial, sans-serif; padding: 20px; }}
                h1 {{ color: #1e40af; }}
                table {{ width: 100%; border-collapse: collapse; margin: 20px 0; }}
                th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
                th {{ background-color: #1e40af; color: white; }}
                .metric {{ background-color: #f3f4f6; padding: 15px; margin: 10px 0; border-radius: 5px; }}
            </style>
        </head>
        <body>
            <h1>Daily Flash Report</h1>
            <p><strong>Date:</strong> {flash_data["date"]}</p>

            <div class="metric">
                <h3>Occupancy</h3>
                <p>Occupied Rooms: {flash_data["occupancy"]["occupied_rooms"]}</p>
                <p>Total Rooms: {flash_data["occupancy"]["total_rooms"]}</p>
                <p>Occupancy %: {flash_data["occupancy"]["occupancy_rate"]:.1f}%</p>
            </div>

            <div class="metric">
                <h3>Revenue</h3>
                <p>Room Revenue: {currency_symbol}{flash_data["revenue"]["room_revenue"]:.2f}</p>
                <p>Total Revenue: {currency_symbol}{flash_data["revenue"]["total_revenue"]:.2f}</p>
                <p>ADR: {currency_symbol}{flash_data["revenue"]["adr"]:.2f}</p>
                <p>RevPAR: {currency_symbol}{flash_data["revenue"]["rev_par"]:.2f}</p>
            </div>

            <div class="metric">
                <h3>Arrivals &amp; Departures</h3>
                <p>Arrivals: {flash_data["movements"]["arrivals"]}</p>
                <p>Departures: {flash_data["movements"]["departures"]}</p>
                <p>Stayovers: {flash_data["movements"]["stayovers"]}</p>
            </div>
        </body>
        </html>
        """

        try:
            from weasyprint import HTML  # type: ignore
        except (ImportError, OSError) as exc:
            raise HTTPException(status_code=503, detail="PDF oluşturucu bu ortamda kullanılamıyor") from exc
        pdf_bytes = HTML(string=html_content).write_pdf()
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="daily-flash-{flash_data["date"]}.pdf"'},
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF generation failed: {str(e)}") from e


@sub_router.post("/reports/email-daily-flash")
async def email_daily_flash(
    data: dict,
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_reports")),  # v85 DU follow-up: POST architect MEDIUM
):
    """
    Email daily flash report to recipients
    """
    recipients = data.get("recipients", [])

    if not recipients:
        raise HTTPException(status_code=400, detail="Recipients list is required")

    try:
        flash_data = await get_daily_flash_report(None, current_user)
        from core.tenant_currency import get_tenant_currency

        _, currency_symbol = await get_tenant_currency(current_user.tenant_id)

        email_html = f"""
        <html>
        <head>
            <style>
                body {{ font-family: Arial, sans-serif; }}
                .metric {{ background-color: #f3f4f6; padding: 15px; margin: 10px 0; border-radius: 5px; }}
                h3 {{ color: #1e40af; }}
            </style>
        </head>
        <body>
            <h2>Daily Flash Report - {flash_data["date"]}</h2>

            <div class="metric">
                <h3>Occupancy</h3>
                <p>Occupied: {flash_data["occupancy"]["occupied_rooms"]} / {flash_data["occupancy"]["total_rooms"]} ({flash_data["occupancy"]["occupancy_rate"]:.1f}%)</p>
            </div>

            <div class="metric">
                <h3>Revenue</h3>
                <p>Room Revenue: {currency_symbol}{flash_data["revenue"]["room_revenue"]:.2f}</p>
                <p>Total Revenue: {currency_symbol}{flash_data["revenue"]["total_revenue"]:.2f}</p>
            </div>

            <div class="metric">
                <h3>Movements</h3>
                <p>Arrivals: {flash_data["movements"]["arrivals"]}</p>
                <p>Departures: {flash_data["movements"]["departures"]}</p>
            </div>

            <p><small>Generated by Syroce PMS</small></p>
        </body>
        </html>
        """

        subject = f'Daily Flash Report - {flash_data["date"]}'
        results = await asyncio.gather(
            *[send_email(to=r, subject=subject, html=email_html) for r in recipients],
            return_exceptions=True,
        )
        sent = [r for r, res in zip(recipients, results, strict=False) if isinstance(res, dict) and res.get("sent")]
        failed = [r for r in recipients if r not in sent]

        logger.info("Daily flash email: %d/%d delivered (failed=%s)", len(sent), len(recipients), failed)

        return {
            "success": len(sent) > 0,
            "message": f"Daily flash report sent to {len(sent)}/{len(recipients)} recipients",
            "recipients_sent": sent,
            "recipients_failed": failed,
            "provider": next(
                (r.get("provider") for r in results if isinstance(r, dict) and r.get("sent")),
                None,
            ),
        }

    except Exception as e:
        logger.exception("email_daily_flash failed")
        raise HTTPException(status_code=500, detail=f"Email sending failed: {str(e)}") from e


@sub_router.get("/reports/daily-flash")
@cached(ttl=15, key_prefix="report_daily_flash")  # Cache for 15 seconds
async def get_daily_flash_report(
    date_str: str | None = None,
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_reports")),  # v85 DU: daily flash GM/CFO
    _nocache: bool = Query(False, alias="nocache"),
):
    """Daily Flash Report - GM/CFO Dashboard"""
    if date_str:
        target_date = _report_date(date_str, "date_str")
    else:
        state = await ensure_business_date_initialized(db, current_user.tenant_id)
        target_date = datetime.fromisoformat(state["business_date"][:10]).date()
    day_key = target_date.isoformat()
    next_key = (target_date + timedelta(days=1)).isoformat()
    metrics = await load_stay_night_metrics(db, current_user.tenant_id, target_date, target_date, actual_only=True)
    metric = metrics[0] if metrics else {"occupied_rooms": 0, "total_rooms": 0, "occupancy_rate": 0}
    total_rooms = metric["total_rooms"]
    occupied_rooms = metric["occupied_rooms"]
    occupancy_rate = metric["occupancy_rate"]
    movement_bookings = await db.bookings.find(
        {"tenant_id": current_user.tenant_id, "$or": [{"check_in": {"$gte": day_key, "$lt": next_key}}, {"check_out": {"$gte": day_key, "$lt": next_key}}]},
        {"_id": 0, "check_in": 1, "check_out": 1, "status": 1},
    ).to_list(10000)
    arrivals = sum(1 for booking in movement_bookings if as_date(booking.get("check_in")) == target_date and str(booking.get("status") or "").lower() not in NON_COMMERCIAL_STATUSES)
    departures = sum(1 for booking in movement_bookings if as_date(booking.get("check_out")) == target_date and str(booking.get("status") or "").lower() not in NON_COMMERCIAL_STATUSES)

    # Note: Revenue is calculated from folio charges, not bookings directly

    # Night-audit charges are posted after the business day closes.  Their
    # timestamp therefore belongs to the next calendar day; report by the
    # accounting business date instead.
    charges = await db.folio_charges.find(
        {
            "tenant_id": current_user.tenant_id,
            "voided": {"$ne": True},
            "$or": [
                {"business_date": day_key},
                {"business_date": {"$exists": False}, "date": {"$regex": f"^{day_key}"}},
                {"business_date": None, "date": {"$regex": f"^{day_key}"}},
            ],
        }
    ).to_list(10000)

    posted_total = sum(float(c.get("total") or c.get("amount") or 0) for c in charges)

    # Revenue breakdown by category
    room_revenue = sum(
        float(c.get("total") or c.get("amount") or 0)
        for c in charges
        if str(c.get("charge_category") or c.get("charge_type") or "").lower() in {"room", "accommodation", "room_charge"}
    )
    posted_fb_revenue = sum(
        float(c.get("total") or c.get("amount") or 0)
        for c in charges
        if str(c.get("charge_category") or c.get("charge_type") or "").lower() in FNB_CATEGORIES
    )
    pos_orders = await db.pos_orders.find(
        {
            "tenant_id": current_user.tenant_id,
            "status": {"$nin": ["cancelled", "canceled", "void", "voided", "refunded"]},
            **accounting_day_match(
                day_key,
                {"closed_at": {"$regex": f"^{day_key}"}},
                {"created_at": {"$regex": f"^{day_key}"}},
            ),
        },
        {"_id": 0, "total_amount": 1, "grand_total": 1},
    ).to_list(5000)
    pos_fb_revenue = sum(float(order.get("grand_total") or order.get("total_amount") or 0) for order in pos_orders)
    fb_revenue = posted_fb_revenue if posted_fb_revenue else pos_fb_revenue
    total_revenue = posted_total if posted_fb_revenue else posted_total + pos_fb_revenue
    other_revenue = max(total_revenue - room_revenue - fb_revenue, 0)

    # Calculate ADR and RevPAR
    adr = round(room_revenue / occupied_rooms, 2) if occupied_rooms > 0 else 0
    rev_par = round(room_revenue / total_rooms, 2) if total_rooms > 0 else 0

    return {
        "date": day_key,
        "occupancy": {"occupied_rooms": occupied_rooms, "total_rooms": total_rooms, "occupancy_rate": occupancy_rate},
        "movements": {"arrivals": arrivals, "departures": departures, "stayovers": max(occupied_rooms - arrivals, 0)},
        "revenue": {
            "total_revenue": round(total_revenue, 2),
            "room_revenue": round(room_revenue, 2),
            "fb_revenue": round(fb_revenue, 2),
            "other_revenue": round(other_revenue, 2),
            "adr": adr,
            "rev_par": rev_par,
            "basis": "posted_folio_charges" if posted_fb_revenue or not pos_fb_revenue else "posted_folio_charges_plus_pos",
        },
    }


@sub_router.get("/reports/daily-flash/excel")
@cached(ttl=600, key_prefix="report_daily_flash_excel")  # Cache for 10 min
async def export_daily_flash_excel(
    date_str: str | None = None,
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_reports")),  # v85 DU: daily flash excel
):
    """Export Daily Flash Report to Excel with professional formatting and dynamic currency"""
    report_data = await get_daily_flash_report(date_str, current_user)
    target_date = report_data["date"]

    # Dynamically fetch currency symbol for the tenant
    from core.tenant_currency import get_tenant_currency
    _, currency_symbol = await get_tenant_currency(current_user.tenant_id)

    # Try using our enhanced openpyxl logic if available, else fallback
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        wb = Workbook()
        ws = wb.active
        ws.title = f"Flash Report {target_date}"

        title_font = Font(size=16, bold=True, color="333333")
        header_font = Font(size=12, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        border_side = Side(border_style="thin", color="CCCCCC")
        box_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

        ws["A1"] = "DAILY FLASH REPORT / GÜNLÜK ÖZET"
        ws["A1"].font = title_font
        ws.merge_cells("A1:B1")

        ws["A2"] = f"Date: {target_date}"
        ws["A2"].font = Font(size=12, italic=True)
        ws.merge_cells("A2:B2")

        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 25

        row_num = 4

        sections = [
            ("OCCUPANCY / DOLULUK", [
                ("Total Rooms", report_data["occupancy"]["total_rooms"], ""),
                ("Occupied Rooms", report_data["occupancy"]["occupied_rooms"], ""),
                ("Occupancy Rate", f"{report_data['occupancy']['occupancy_rate']}%", "")
            ]),
            ("MOVEMENTS / HAREKETLER", [
                ("Arrivals (Check-in)", report_data["movements"]["arrivals"], ""),
                ("Departures (Check-out)", report_data["movements"]["departures"], ""),
                ("Stayovers", report_data["movements"]["stayovers"], "")
            ]),
            ("REVENUE / GELİRLER", [
                ("Total Revenue", report_data['revenue']['total_revenue'], currency_symbol),
                ("Room Revenue", report_data['revenue']['room_revenue'], currency_symbol),
                ("F&B Revenue", report_data['revenue']['fb_revenue'], currency_symbol),
                ("Other Revenue", report_data['revenue']['other_revenue'], currency_symbol),
                ("ADR (Average Daily Rate)", report_data['revenue']['adr'], currency_symbol),
                ("RevPAR (Revenue Per Available Room)", report_data['revenue']['rev_par'], currency_symbol)
            ])
        ]

        for section_title, items in sections:
            ws.cell(row=row_num, column=1, value=section_title).font = header_font
            ws.cell(row=row_num, column=1).fill = header_fill
            ws.cell(row=row_num, column=2, value="").fill = header_fill
            ws.cell(row=row_num, column=1).alignment = Alignment(horizontal="left")
            row_num += 1

            for label, value, cur in items:
                ws.cell(row=row_num, column=1, value=label).border = box_border
                val_cell = ws.cell(row=row_num, column=2, value=f"{cur}{value:,.2f}" if cur else value)
                val_cell.border = box_border
                val_cell.alignment = Alignment(horizontal="right")
                row_num += 1
            row_num += 1

        import tempfile

        from fastapi.responses import FileResponse
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(path)
        return FileResponse(path, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=f"daily_flash_report_{target_date}.xlsx")

    except ImportError:
        # Fallback to simple
        headers = ["Metric", "Value"]
        data = [
            ["Report Date", target_date],
            ["", ""],
            ["OCCUPANCY", ""],
            ["Total Rooms", report_data["occupancy"]["total_rooms"]],
            ["Occupied Rooms", report_data["occupancy"]["occupied_rooms"]],
            ["Occupancy Rate", f"{report_data['occupancy']['occupancy_rate']}%"],
            ["", ""],
            ["MOVEMENTS", ""],
            ["Arrivals", report_data["movements"]["arrivals"]],
            ["Departures", report_data["movements"]["departures"]],
            ["Stayovers", report_data["movements"]["stayovers"]],
            ["", ""],
            ["REVENUE", ""],
            ["Total Revenue", f"{currency_symbol}{report_data['revenue']['total_revenue']:,.2f}"],
            ["Room Revenue", f"{currency_symbol}{report_data['revenue']['room_revenue']:,.2f}"],
            ["F&B Revenue", f"{currency_symbol}{report_data['revenue']['fb_revenue']:,.2f}"],
            ["Other Revenue", f"{currency_symbol}{report_data['revenue']['other_revenue']:,.2f}"],
            ["ADR (Average Daily Rate)", f"{currency_symbol}{report_data['revenue']['adr']:,.2f}"],
            ["RevPAR", f"{currency_symbol}{report_data['revenue']['rev_par']:,.2f}"],
        ]
        wb = create_excel_workbook(title=f"Daily Flash Report - {target_date}", headers=headers, data=data, sheet_name="Daily Flash")
        filename = f"daily_flash_report_{target_date}.xlsx"
        return excel_response(wb, filename)


@sub_router.post("/reports/send-weekly-email")
async def send_weekly_management_email(
    email_config: dict,
    credentials: HTTPAuthorizationCredentials = Depends(security),
    _perm=Depends(require_op("view_reports")),  # v92 DW
):
    """Send weekly management summary via email"""
    current_user = await get_current_user(credentials)

    state = await ensure_business_date_initialized(db, current_user.tenant_id)
    week_end = datetime.fromisoformat(state["business_date"][:10]).date()
    week_start = week_end - timedelta(days=6)
    start_key = week_start.isoformat()
    end_key = week_end.isoformat()
    next_key = (week_end + timedelta(days=1)).isoformat()

    total_bookings = await db.bookings.count_documents(
        {
            "tenant_id": current_user.tenant_id,
            "$or": [
                {"created_at": {"$gte": start_key, "$lt": next_key}},
                {"reservation_date": {"$gte": start_key, "$lt": next_key}},
            ],
            "status": {"$nin": list(NON_COMMERCIAL_STATUSES)},
        }
    )
    charges = await db.folio_charges.find(
        {
            "tenant_id": current_user.tenant_id,
            "voided": {"$ne": True},
            **accounting_period_match(
                start_key,
                end_key,
                {"date": {"$gte": start_key, "$lt": next_key}},
                {"created_at": {"$gte": start_key, "$lt": next_key}},
            ),
        },
        {"_id": 0, "total": 1, "amount": 1, "charge_category": 1, "charge_type": 1},
    ).to_list(20000)
    total_revenue = sum(float(charge.get("total") or charge.get("amount") or 0) for charge in charges)
    room_revenue = sum(
        float(charge.get("total") or charge.get("amount") or 0)
        for charge in charges
        if str(charge.get("charge_category") or charge.get("charge_type") or "").lower() in {"room", "accommodation", "room_charge"}
    )
    stay_metrics = await load_stay_night_metrics(db, current_user.tenant_id, week_start, week_end, actual_only=True)
    occupied_room_nights = sum(int(item.get("occupied_rooms") or 0) for item in stay_metrics)
    available_room_nights = sum(int(item.get("total_rooms") or 0) for item in stay_metrics)
    occupancy = round(occupied_room_nights / available_room_nights * 100, 2) if available_room_nights else 0.0
    adr = round(room_revenue / occupied_room_nights, 2) if occupied_room_nights else 0.0
    revpar = round(room_revenue / available_room_nights, 2) if available_room_nights else 0.0

    recipient = str(email_config.get("email") or current_user.email or "").strip()
    if not recipient:
        raise HTTPException(status_code=422, detail="Rapor alıcısı e-posta adresi gerekli")
    subject = f"Haftalık Yönetim Özeti - {start_key} / {end_key}"
    report_data = {
        "week_start": start_key,
        "week_ending": end_key,
        "total_bookings": total_bookings,
        "total_revenue": round(total_revenue, 2),
        "room_revenue": round(room_revenue, 2),
        "key_metrics": {"occupancy": occupancy, "adr": adr, "revpar": revpar},
    }
    html = f"""
    <h2>Haftalık Yönetim Özeti</h2>
    <p><strong>Dönem:</strong> {start_key} – {end_key}</p>
    <ul>
      <li>Yeni rezervasyon: {total_bookings}</li>
      <li>Toplam gelir: {total_revenue:,.2f}</li>
      <li>Oda geliri: {room_revenue:,.2f}</li>
      <li>Doluluk: %{occupancy:.2f}</li>
      <li>ADR: {adr:,.2f}</li>
      <li>RevPAR: {revpar:,.2f}</li>
    </ul>
    """
    delivery = await send_email(to=recipient, subject=subject, html=html)
    delivered = isinstance(delivery, dict) and bool(delivery.get("sent"))

    email_record = {
        "id": str(uuid.uuid4()),
        "tenant_id": current_user.tenant_id,
        "recipient_email": recipient,
        "subject": subject,
        "report_type": "weekly_summary",
        "report_data": report_data,
        "status": "sent" if delivered else "failed",
        "sent_at": datetime.now(UTC).isoformat() if delivered else None,
        "sent_by": current_user.name,
        "provider": delivery.get("provider") if isinstance(delivery, dict) else None,
    }

    await db.email_reports.insert_one(email_record)

    if not delivered:
        raise HTTPException(status_code=502, detail="Haftalık rapor e-postası gönderilemedi")
    return {"message": "Haftalık yönetim özeti gönderildi", "email_id": email_record["id"], "recipient": recipient}


@sub_router.get("/reports/email-history")
async def get_email_report_history(
    limit: int = Query(20, ge=1, le=100),
    credentials: HTTPAuthorizationCredentials = Depends(security),
    _perm=Depends(require_op("view_reports")),
):
    """Get email report history"""
    current_user = await get_current_user(credentials)

    emails = []
    async for email in db.email_reports.find({"tenant_id": current_user.tenant_id}).sort("sent_at", -1).limit(limit):
        email.pop("_id", None)
        emails.append(email)

    return {"emails": emails, "count": len(emails)}


@sub_router.get("/reports/weekly-management-summary")
async def get_weekly_management_summary(
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_reports")),
):
    """Get weekly management summary report"""
    state = await ensure_business_date_initialized(db, current_user.tenant_id)
    week_end = datetime.fromisoformat(str(state["business_date"])[:10]).date()
    week_start = week_end - timedelta(days=6)
    end_exclusive = week_end + timedelta(days=1)

    # Get key metrics for the week
    total_bookings = await db.bookings.count_documents(
        {
            "tenant_id": current_user.tenant_id,
            "created_at": {"$gte": week_start.isoformat(), "$lt": end_exclusive.isoformat()},
            "status": {"$nin": list(NON_COMMERCIAL_STATUSES)},
        }
    )
    charges = await db.folio_charges.find(
        {
            "tenant_id": current_user.tenant_id,
            "voided": {"$ne": True},
            "$or": [
                {"business_date": {"$gte": week_start.isoformat(), "$lte": week_end.isoformat()}},
                {"business_date": {"$exists": False}, "date": {"$gte": week_start.isoformat(), "$lt": end_exclusive.isoformat()}},
                {"business_date": None, "date": {"$gte": week_start.isoformat(), "$lt": end_exclusive.isoformat()}},
            ],
        },
        {"_id": 0, "total": 1, "amount": 1},
    ).to_list(50000)
    total_revenue = sum(float(charge.get("total") or charge.get("amount") or 0) for charge in charges)
    metrics = await load_stay_night_metrics(db, current_user.tenant_id, week_start, week_end, actual_only=True)
    occupied_room_nights = sum(int(row.get("occupied_rooms") or 0) for row in metrics)
    available_room_nights = sum(int(row.get("total_rooms") or 0) for row in metrics)
    occupied_avg = occupied_room_nights / available_room_nights * 100 if available_room_nights else 0

    # Get maintenance tasks completed
    completed_tasks = await db.maintenance_tasks.count_documents(
        {
            "tenant_id": current_user.tenant_id,
            "status": "completed",
            "completed_at": {"$gte": week_start.isoformat(), "$lt": end_exclusive.isoformat()},
        }
    )

    # B: gerçek misafir memnuniyeti (haftalık review rating ortalaması); review yoksa fail-closed (null)
    sat_docs = await db.reviews.find(
        {"tenant_id": current_user.tenant_id, "created_at": {"$gte": week_start.isoformat(), "$lt": end_exclusive.isoformat()}},
        {"_id": 0, "rating": 1},
    ).to_list(2000)
    sat_ratings = [r.get("rating") for r in sat_docs if isinstance(r.get("rating"), (int, float))]
    guest_satisfaction = round(sum(sat_ratings) / len(sat_ratings), 2) if sat_ratings else None

    return {
        "week_start": week_start.isoformat(),
        "week_ending": week_end.isoformat(),
        "total_bookings": total_bookings,
        "total_revenue": round(total_revenue, 2),
        "revenue_basis": "posted_folio_charges",
        "avg_occupancy_pct": round(occupied_avg, 2),
        "completed_maintenance": completed_tasks,
        "guest_satisfaction": guest_satisfaction,
        "guest_satisfaction_available": guest_satisfaction is not None,
        "guest_satisfaction_reviews": len(sat_ratings),
        # top_performers: gerçek haftalık personel-performans kaynağı yok -> fail-closed (sahte personel kaldırıldı)
        "top_performers": [],
        "top_performers_available": False,
    }
