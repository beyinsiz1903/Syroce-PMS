"""Reservation analysis and market segments use the same booked-sale contract."""

import io
from datetime import UTC, date, datetime
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from routers.departments import reports as segments
from routers.reports_pkg import reservation_performance as reports

pytest_plugins = ("tests.test_report_ledger_mongo_acceptance",)
START, END = date(2026, 9, 1), date(2026, 9, 30)
USER = SimpleNamespace(tenant_id="A")


def booking(**extra):
    return {"tenant_id": "A", "id": "b", "check_in": "2026-09-10", "check_out": "2026-09-12",
        "created_at": "2026-09-01", "status": "confirmed", "total_amount": 100, "currency": "EUR", **extra}


def aggregate(*rows, **options):
    return reports.aggregate_reservation_performance(list(rows), start_date=START, end_date=END, **options)


@pytest.mark.parametrize("status", ["draft", "pending", "tentative", "unknown", "", None])
def test_unconfirmed_statuses_remain_visible_but_never_become_booked_sales(status):
    result = aggregate(booking(status=status))
    assert result["summary"]["total_bookings"] == 1
    assert result["summary"]["commercial_bookings"] == 0
    assert result["summary"]["booked_revenue"] == 0
    assert result["summary"]["total_room_nights"] == 0
    assert result["daily_arrivals"][0]["commercial_reservations"] == 0


def test_comp_reference_value_is_zero_in_summary_channels_daily_and_detail():
    result = aggregate(booking(is_complimentary=True, total_amount=900))
    assert result["summary"]["commercial_bookings"] == 1
    assert result["summary"]["total_room_nights"] == 2
    assert result["summary"]["booked_revenue_by_currency"] == {"EUR": 0}
    assert result["channel_breakdown"][0]["revenue"] == 0
    assert result["daily_arrivals"][0]["revenue"] == 0
    assert result["rows"][0]["total_amount"] == 0


@pytest.mark.parametrize("extra,expected", [
    ({"total_amount": None, "total": 123.45}, 123.45),
    ({"total_amount": 0, "total": 999}, 0),
    ({"total_amount": None, "rate_per_night": 12.35}, 24.70),
    ({"total_amount": "1.005"}, 1.01),
])
def test_legacy_amount_precedence_explicit_zero_and_half_up(extra, expected):
    result = aggregate(booking(**extra))
    assert result["summary"]["booked_revenue"] == expected
    assert result["rows"][0]["total_amount"] == expected


@pytest.mark.parametrize("value", ["NaN", "Infinity", "bad", True])
def test_invalid_amount_is_explicit_not_a_successful_zero_sale(value):
    result = aggregate(booking(total_amount=value))
    assert result["rows"][0]["total_amount"] is None
    assert result["data_quality"]["invalid_amounts"] == 1
    assert result["summary"]["booked_revenue"] == 0


@pytest.mark.parametrize("check_out", ["bad", "2026-09-10", "2026-09-09", None])
def test_invalid_stays_do_not_inflate_sales_or_average_stay(check_out):
    result = aggregate(booking(check_out=check_out))
    assert result["summary"]["commercial_bookings"] == 0
    assert result["summary"]["booked_revenue"] == 0
    assert result["data_quality"]["invalid_stays"] == 1


def test_future_creation_is_not_turned_into_same_day_lead_time():
    result = aggregate(booking(created_at="2026-09-11"))
    assert result["rows"][0]["lead_time_days"] is None
    assert sum(row["count"] for row in result["lead_time_breakdown"]) == 0
    assert result["data_quality"]["invalid_lead_times"] == 1


def test_missing_booking_currency_follows_property_not_try():
    result = aggregate(booking(currency=None), fallback_currency="GBP")
    assert result["summary"]["booked_revenue_by_currency"] == {"GBP": 100}
    assert result["currency"] == result["rows"][0]["currency"] == "GBP"


@pytest.mark.parametrize("day", ["20260901", "2026-09-01T00:00:00", "2026-02-30"])
async def test_noncanonical_report_dates_fail_validation(day):
    with pytest.raises(HTTPException) as invalid:
        await reports._resolve_report_dates(tenant_id="A", start_date=day, end_date="2026-09-30")
    assert invalid.value.status_code == 422


@pytest.mark.parametrize("check_in", ["2026-09-01", datetime(2026, 9, 30, 23, 59, tzinfo=UTC)])
async def test_real_query_preserves_projection_currency_and_all_rows(ledger_db, monkeypatch, check_in):
    await ledger_db.bookings.insert_many([
        booking(id="comp", check_in=check_in, check_out="2026-10-02", is_complimentary=True, currency=None),
        booking(id="legacy", check_in=check_in, check_out="2026-10-02", total_amount=None, total=55.55, currency=None),
        booking(id="foreign", tenant_id="B", total_amount=999999),
        booking(id="next", check_in="2026-10-01", check_out="2026-10-02", total_amount=999999),
    ])
    monkeypatch.setattr(reports, "db", ledger_db)
    monkeypatch.setattr(reports, "get_tenant_currency", AsyncMock(return_value=("GBP", "£")), raising=False)
    result = await reports._reservation_performance_payload(tenant_id="A", start_date=START, end_date=END)
    assert result["summary"]["total_bookings"] == 2
    assert result["summary"]["booked_revenue_by_currency"] == {"GBP": 55.55}
    assert {row["booking_id"] for row in result["rows"]} == {"comp", "legacy"}


async def test_market_segment_excludes_unconfirmed_and_uses_paid_currency_nights(ledger_db, monkeypatch):
    await ledger_db.bookings.insert_many([
        booking(id="paid", total_amount=100, market_segment="direct"),
        booking(id="comp", total_amount=999, is_complimentary=True, market_segment="direct"),
        booking(id="usd", total_amount=200, currency="USD", market_segment="direct"),
        *[booking(id=status, status=status, total_amount=999999) for status in ["pending", "draft", "tentative"]],
        booking(id="foreign", tenant_id="B", total_amount=999999),
    ])
    monkeypatch.setattr(segments, "db", ledger_db)
    monkeypatch.setattr(segments, "_native_report_currency", AsyncMock(return_value="EUR"))
    result = await segments.get_market_segment_report.__wrapped__("2026-09-01", "2026-09-30", USER)
    assert result["total_bookings"] == 3
    row = result["market_segments"]["direct"]
    assert row["nights"] == 6
    assert row["revenue_by_currency"] == {"EUR": 100, "USD": 200}
    assert row["adr_by_currency"] == {"EUR": 50, "USD": 100}
    assert row["sold_nights_by_currency"] == {"EUR": 2, "USD": 2}


async def test_excel_values_agree_with_summary_and_detail(ledger_db, monkeypatch):
    await ledger_db.bookings.insert_many([booking(id="comp", is_complimentary=True, total_amount=900),
        booking(id="eur", total_amount="1.005"), booking(id="usd", currency="USD", total_amount=22.33)])
    monkeypatch.setattr(reports, "db", ledger_db)
    monkeypatch.setattr(reports, "get_tenant_currency", AsyncMock(return_value=("EUR", "€")), raising=False)
    response = await reports.export_reservation_performance_excel("2026-09-01", "2026-09-30", USER)
    content = b"".join([part async for part in response.body_iterator])
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    summary = dict(workbook["Özet"].iter_rows(min_row=3, values_only=True))
    assert "1.01 EUR" in summary["Geçerli rezervasyon bedeli"]
    assert "22.33 USD" in summary["Geçerli rezervasyon bedeli"]
    amounts = [row[7] for row in workbook["Rezervasyonlar"].iter_rows(min_row=2, values_only=True)]
    assert sorted(amounts) == [0, 1.01, 22.33]


async def test_market_segment_excel_preserves_unknown_money_and_paid_denominators(ledger_db, monkeypatch):
    await ledger_db.bookings.insert_many([
        booking(total_amount=100, market_segment="=bad", rate_type="@bad"),
        booking(id="comp", is_complimentary=True, total_amount=900, market_segment="=bad"),
        booking(id="bad", currency="USD", total_amount="NaN", market_segment="=bad"),
    ])
    monkeypatch.setattr(segments, "db", ledger_db)
    monkeypatch.setattr(segments, "_native_report_currency", AsyncMock(return_value="EUR"))
    response = await segments.export_market_segment_excel("2026-09-01", "2026-09-30", USER)
    workbook = load_workbook(io.BytesIO(b"".join([part async for part in response.body_iterator])), data_only=False)
    rows = list(workbook["Market Segments"].iter_rows(min_row=3, values_only=True))
    eur, usd = sorted(rows, key=lambda row: row[1])
    assert eur[0].startswith("'") and eur[1:] == ("EUR", 2, 4, 100, 50, 2, 0)
    assert usd[1:] == ("USD", 1, 2, None, None, 0, 1)
    assert workbook["Notes"]["B4"].value == 1


@pytest.mark.parametrize("first,last", [("20260901", "2026-09-30"), ("2026-09-01", "9999-12-31"),
    ("2026-09-01T12:00:00", "2026-09-30"), ("2020-01-01", "2026-09-30")])
async def test_market_dates_are_bounded_and_canonical(first, last):
    with pytest.raises(HTTPException) as invalid:
        await segments.get_market_segment_report.__wrapped__(first, last, USER)
    assert invalid.value.status_code == 422


def test_decimal128_and_rounded_details_reconcile_with_channel_and_total():
    from bson.decimal128 import Decimal128

    result = aggregate(*[booking(id=str(index), total_amount=Decimal128("0.005")) for index in range(3)])
    assert result["summary"]["booked_revenue"] == 0.03
    assert result["channel_breakdown"][0]["revenue"] == 0.03
    assert result["daily_arrivals"][0]["revenue"] == 0.03
    assert sum(row["total_amount"] for row in result["rows"]) == 0.03


async def test_real_query_does_not_truncate_large_reports(ledger_db, monkeypatch):
    await ledger_db.bookings.insert_many([booking(id=str(index), total_amount=0.01) for index in range(1005)])
    monkeypatch.setattr(reports, "db", ledger_db)
    monkeypatch.setattr(reports, "get_tenant_currency", AsyncMock(return_value=("EUR", "€")))
    result = await reports._reservation_performance_payload(tenant_id="A", start_date=START, end_date=END)
    assert len(result["rows"]) == result["summary"]["total_bookings"] == 1005
    assert result["summary"]["booked_revenue_by_currency"] == {"EUR": 10.05}


@pytest.mark.parametrize("report", ["reservation", "market"])
async def test_deploy_does_not_reuse_cached_pre_fix_calculations(monkeypatch, report):
    from cache_manager import cache

    old_prefix = "report_reservation_performance" if report == "reservation" else "report_market_segment_v3"
    keys = []

    def get(key):
        keys.append(key)
        return {"stale_pre_fix_calculation": True} if f":{old_prefix}:" in key else None

    monkeypatch.setattr(cache, "enabled", True)
    monkeypatch.setattr(cache, "get", get)
    monkeypatch.setattr(cache, "set", lambda *_args, **_kwargs: True)
    monkeypatch.setattr(reports, "_reservation_performance_payload", AsyncMock(return_value={"fresh": True}))
    monkeypatch.setattr(segments, "db", SimpleNamespace(bookings=SimpleNamespace(find=lambda *_args: SimpleNamespace(to_list=AsyncMock(return_value=[])))))
    monkeypatch.setattr(segments, "_native_report_currency", AsyncMock(return_value="EUR"))
    endpoint = reports.get_reservation_performance_report if report == "reservation" else segments.get_market_segment_report
    result = await endpoint(start_date="2026-09-01", end_date="2026-09-30", current_user=USER, _nocache=False)
    assert not result.get("stale_pre_fix_calculation")
    assert keys and all(f":{old_prefix}:" not in key for key in keys)
