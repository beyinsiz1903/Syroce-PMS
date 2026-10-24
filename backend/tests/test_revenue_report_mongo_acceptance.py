from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from routers.finance import folio

pytest_plugins = ("tests.test_report_ledger_mongo_acceptance",)
USER = SimpleNamespace(tenant_id="A")
DAY = "2026-09-05"


async def test_category_endpoint_shares_posted_pos_refund_tax_room_type_and_export_snapshot(ledger_db, monkeypatch):
    db = ledger_db
    monkeypatch.setattr(folio, "db", db)
    tid = "b3c7bb47-1234-1234-1234-123456789012"
    await db.room_types.insert_one({"tenant_id": "A", "id": tid, "name": "Standard"})
    await db.rooms.insert_one({"tenant_id": "A", "id": "r1", "room_type": tid})
    await db.bookings.insert_one({"tenant_id": "A", "id": "b1", "room_id": "r1", "currency": "TRY"})
    await db.folio_charges.insert_many([
        {"tenant_id": "A", "id": "room", "booking_id": "b1", "charge_category": "room", "business_date": DAY, "amount": 100, "tax_amount": 12, "tax_breakdown": {"vat": 10, "accommodation_tax": 2}, "total": 112, "currency": "TRY"},
        {"tenant_id": "A", "id": "zero", "business_date": DAY, "amount": 1, "total": 0, "currency": "TRY"},
        {"tenant_id": "B", "business_date": DAY, "total": 999999, "currency": "TRY"},
    ])
    await db.extra_charges.insert_one({"tenant_id": "A", "id": "transfer", "category": "food", "business_date": DAY, "total": 50, "currency": "TRY", "source_pos_order_id": "transfer"})
    await db.pos_orders.insert_many([
        {"tenant_id": "A", "id": "fallback", "status": "paid", "business_date": DAY, "total_amount": 100, "grand_total": 120, "currency": "TRY"},
        {"tenant_id": "A", "id": "transfer", "status": "paid", "business_date": DAY, "grand_total": 50, "currency": "TRY"},
        {"tenant_id": "A", "id": "served", "status": "served", "business_date": DAY, "grand_total": 9999, "currency": "TRY"},
    ])
    await db.pos_transactions.insert_one({"tenant_id": "A", "id": "refund", "order_id": "old", "status": "refund", "payment_type": "refund", "business_date": DAY, "total_amount": 22, "tip_amount": 2, "currency": "TRY"})
    data = await folio.revenue_by_category.__wrapped__(DAY, DAY, USER, None)
    assert data["totals_by_currency"]["TRY"]["total"] == 262  # 112 + 0 + 50 + 120 − 20
    assert data["revenue_trend"][0]["revenue_by_currency"] == {"TRY": 262}
    assert data["room_type_revenue"] == [{"name": "Standard", "revenue_by_currency": {"TRY": 112}}]
    room = next(row for row in data["rows"] if row["category"] == "room")
    assert (room["net"], room["vat"], room["city_tax"]) == (100, 10, 2)
    assert data["reconciliation"]["issue_count"] == 1
    result = await folio.revenue_by_category_excel(DAY, DAY, data["snapshot"], USER, None)
    blob = b"".join([chunk async for chunk in result.body_iterator])
    book = load_workbook(BytesIO(blob), data_only=True)
    totals = [row for row in book.active.iter_rows(values_only=True) if row[0] == "TOPLAM"]
    assert totals[0][-1] == 262
    assert await db.folio_charges.count_documents({"tenant_id": "A"}) == 2  # report never repairs the ledger
    await db.folio_charges.insert_one({"tenant_id": "A", "business_date": DAY, "total": 1, "currency": "TRY"})
    with pytest.raises(HTTPException) as error:
        await folio.revenue_by_category_excel(DAY, DAY, data["snapshot"], USER, None)
    assert error.value.status_code == 409


async def test_forecast_matrix_uses_same_group_blocks_catalog_and_denominator(ledger_db, monkeypatch):
    from routers.reports_pkg import frontdesk_forecast as api

    db = ledger_db
    monkeypatch.setattr(api, "db", db)
    monkeypatch.setattr(api, "ensure_business_date_initialized", AsyncMock(return_value={"business_date": DAY}))
    monkeypatch.setattr(api, "get_tenant_currency", AsyncMock(return_value=("TRY", "₺")))
    type_id = "b3c7bb47-1234-1234-1234-123456789012"
    await db.room_types.insert_one({"tenant_id": "A", "id": type_id, "name": "Standard"})
    await db.rooms.insert_many([{"tenant_id": "A", "id": f"r{i}", "room_type": kind} for i, kind in enumerate([type_id, "standard", "Standard", "Suite"])])
    await db.bookings.insert_one({"tenant_id": "A", "id": "b1", "room_id": "r0", "room_type": type_id, "check_in": DAY, "check_out": "2026-09-06", "status": "confirmed", "adults": 2, "children": 0, "currency": "TRY", "total_amount": 112, "group_block_id": "g1"})
    await db.group_blocks.insert_one({"tenant_id": "A", "id": "g1", "check_in": DAY, "check_out": "2026-09-06", "total_rooms": 2, "room_breakdown": {type_id: 2}, "status": "definite"})
    data = await api.get_frontdesk_forecast(DAY, DAY, current_user=USER)
    assert data["rows"][0]["total_occupied_rooms"] == 2
    groups = {row["label"]: row for row in data["room_type_groups"]}
    assert groups["Standard"]["rows"][0]["total_occupied_rooms"] == 2
    assert groups["Standard"]["rows"][0]["remaining_rooms"] == 1
    assert groups["Standard"]["rows"][0]["occupancy_pct"] == 66.67
    assert sum(group["rows"][0]["total_occupied_rooms"] for group in groups.values()) == data["rows"][0]["total_occupied_rooms"]
    response = await api.get_frontdesk_forecast(DAY, DAY, current_user=USER, output="xlsx", snapshot=data["snapshot"], detail="summary")
    book = load_workbook(BytesIO(b"".join([chunk async for chunk in response.body_iterator])), data_only=True)
    assert book.active.max_column == 7  # date + six summary columns, not twenty
    assert book["Oda Tipi Matrisi"].cell(4, 2).value == "66.67% / +1"


async def test_flash_distinguishes_pending_room_prices_from_posted_revenue(ledger_db, monkeypatch):
    from routers.reports_pkg import flash_email

    monkeypatch.setattr(flash_email, "db", ledger_db)
    await ledger_db.rooms.insert_one({"tenant_id": "A", "id": "r1", "room_type": "Standard"})
    await ledger_db.bookings.insert_one({"tenant_id": "A", "id": "b1", "room_id": "r1", "check_in": DAY, "check_out": "2026-09-06", "status": "checked_in", "currency": "TRY", "total_amount": 5500})
    await ledger_db.folio_charges.insert_one({"tenant_id": "A", "business_date": DAY, "category": "mini_bar", "total": 20, "currency": "TRY"})
    data = await flash_email.get_daily_flash_report.__wrapped__(DAY, USER)
    assert data["revenue"]["adr_by_currency"] == {}
    assert data["revenue"]["posting_pending"] is True
    assert data["revenue"]["operational_adr_by_currency"] == {"TRY": 5500}
    assert data["revenue"]["category_revenue_by_currency"]["minibar"] == {"TRY": 20}
    assert data["revenue"]["fb_revenue_by_currency"] == {}
