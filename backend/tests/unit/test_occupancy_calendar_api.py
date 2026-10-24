import os
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

os.environ.setdefault("JWT_SECRET", "unit-test-secret-key-at-least-32-chars!!")

from routers.reports_pkg import occupancy_calendar as api


@pytest.mark.parametrize('first,last', [('2026-1-01', '2026-01-02'), ('2026-02-30', '2026-03-01'),
                                        ('2026-11-11', '2026-11-10'), ('2026-01-01', '2027-01-02'), ('bad', '2026-01-01')])
def test_range_rejected_before_database(first, last):
    with pytest.raises(HTTPException) as exc:
        api.date_range(first, last)
    assert exc.value.status_code == 422


@pytest.fixture
def dataset(monkeypatch):
    collections = {}
    for name, rows in {'rooms': [{'id': 'r1', 'room_type': 'Standard'}],
                       'bookings': [{'id': 'b', 'status': 'confirmed', 'room_type': 'Standard', 'check_in': '2026-11-10', 'check_out': '2026-11-11'}], 'room_blocks': []}.items():
        collections[name] = SimpleNamespace(find=Mock(return_value=SimpleNamespace(to_list=AsyncMock(return_value=rows))))
    database = SimpleNamespace(**collections)
    monkeypatch.setattr(api, 'db', database)
    monkeypatch.setattr(api, 'ensure_business_date_initialized', AsyncMock(return_value={'business_date': '2026-11-10'}))
    return database


async def read(**kwargs):
    return await api.get_occupancy_calendar('2026-11-10', '2026-11-10', room_type=None, output=kwargs.pop('output', 'json'),
                                          snapshot=kwargs.pop('snapshot', None), current_user=SimpleNamespace(tenant_id='tenant-a'), **kwargs)


@pytest.mark.asyncio
async def test_queries_are_tenant_scoped_uncapped_and_no_pii(dataset):
    result = await read()
    assert result['rows'][0]['expected_rooms'] == 1
    for collection in vars(dataset).values():
        query, projection = collection.find.call_args.args
        assert query['tenant_id'] == 'tenant-a'
        assert not any('guest' in field or 'email' in field for field in projection)
        collection.find.return_value.to_list.assert_awaited_once_with(None)
    assert dataset.rooms.find.call_args.args[1]['is_virtual'] == 1


@pytest.mark.asyncio
async def test_excel_contains_the_same_values_and_refuses_a_changed_snapshot(dataset):
    result = await read()
    response = await read(output='xlsx', snapshot=result['snapshot'])
    raw = b''.join([part async for part in response.body_iterator])
    workbook = load_workbook(BytesIO(raw))
    assert workbook.active.title == 'Oda Tipi Doluluk'
    sheet = workbook['Doluluk Özeti']
    assert sheet.title == 'Doluluk Özeti'
    assert sheet.cell(7, 3).value == result['rows'][0]['expected_rooms']
    assert sheet.cell(7, 2).value == result['rows'][0]['expected_occupancy'] / 100
    assert sheet.cell(7, 2).number_format == '0.00%'
    detail = workbook['Detaylar']
    assert detail.cell(7, 9).value == result['rows'][0]['expected_rooms']
    assert detail.cell(7, 10).value == result['rows'][0]['expected_occupancy'] / 100
    with pytest.raises(HTTPException) as exc:
        await read(output='xlsx', snapshot='stale')
    assert exc.value.status_code == 409


@pytest.mark.asyncio
async def test_one_snapshot_covers_room_types_and_selected_excel_view(dataset):
    result = await read()
    assert result['room_type_groups'][0]['rows'] == result['rows']
    response = await read(output='xlsx', snapshot=result['snapshot'], view='daily', metric='remaining')
    workbook = load_workbook(BytesIO(b''.join([part async for part in response.body_iterator])))
    assert workbook.active.title == 'Doluluk Özeti'
    dataset.bookings.find.return_value.to_list.return_value = []
    with pytest.raises(HTTPException) as exc:
        await read(output='xlsx', snapshot=result['snapshot'])
    assert exc.value.status_code == 409
