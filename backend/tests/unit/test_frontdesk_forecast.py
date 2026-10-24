from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from modules.pms_core.frontdesk_forecast import frontdesk_forecast
from modules.pms_core.frontdesk_forecast_export import forecast_workbook

DAY = date(2026, 10, 24)
ROOMS = [{'id': f'r{i}', 'room_type': 'Standard', 'room_number': str(100 + i)} for i in range(60)]


def booking(key='b', **changes):
    return {'id': key, 'room_type': 'Standard', 'status': 'confirmed', 'check_in': '2026-10-24', 'check_out': '2026-10-25',
        'adults': 2, 'children': 0, 'children_ages': [], 'total_amount': 100, 'currency': 'TRY', **changes}


def report(bookings=(), groups=(), rates=(), closed=(), **kwargs):
    return frontdesk_forecast(ROOMS, list(bookings), list(closed), list(groups), list(rates),
        kwargs.pop('first', DAY), kwargs.pop('last', DAY), kwargs.pop('business_day', DAY), **kwargs)


def test_reference_physical_denominator_and_distinct_closed_capacity():
    bookings = [booking(f'b{i}', room_id=f'r{i}') for i in range(53)]
    closed = [{'room_id': f'r{i}', 'start_date': '2026-10-24', 'end_date': '2026-10-25'} for i in [58, 59]]
    data = report(bookings, closed=closed)
    row = data['rows'][0]
    assert (row['physical_rooms'], row['blocked_rooms'], row['total_occupied_rooms'], row['vacant_rooms']) == (60, 2, 53, 5)
    assert row['occupancy_pct'] == 88.33 and row['sellable_occupancy_pct'] == 91.38
    assert row['booked_rooms'] + row['block_rooms'] + row['vacant_rooms'] + row['blocked_rooms'] == 60
    assert data['totals']['room_revenue_by_currency'] == {'TRY': 5300}


def test_group_block_uses_each_nights_actual_pickup_not_stale_counter():
    group = {'id': 'g', 'room_type': 'Standard', 'total_rooms': 8, 'rooms_picked_up': 999,
        'status': 'definite', 'check_in': '2026-10-24', 'check_out': '2026-10-26'}
    bookings = [booking('g1', group_block_id='g'), booking('g2', group_block_id='g', check_out='2026-10-26'), booking('individual')]
    data = report(bookings, groups=[group], last=date(2026, 10, 25))
    assert [(r['booked_rooms'], r['block_rooms'], r['total_occupied_rooms']) for r in data['rows']] == [(3, 6, 9), (1, 7, 8)]
    assert data['rows'][0]['group_rooms'] == 2
    assert data['rows'][0]['room_revenue_by_currency'] == {'TRY': 250}


def test_room_type_group_breakdown_reconciles_without_double_count():
    rooms = [*ROOMS, {'id': 'suite', 'room_type': 'Suite'}]
    group = {'id': 'g', 'total_rooms': 8, 'room_breakdown': {'Standard': 5, 'Suite': 3},
        'status': 'definite', 'check_in': '2026-10-24', 'check_out': '2026-10-25'}
    bookings = [booking('standard', group_block_id='g'), booking('suite', group_block_id='g', room_id='suite')]
    results = [frontdesk_forecast(rooms, bookings, [], [group], [], DAY, DAY, DAY, room_type=kind) for kind in [None, 'Standard', 'Suite']]
    assert results[0]['rows'][0]['total_occupied_rooms'] == 8
    assert [r['rows'][0]['block_rooms'] for r in results] == [6, 4, 2]
    assert sum(r['rows'][0]['total_occupied_rooms'] for r in results[1:]) == results[0]['rows'][0]['total_occupied_rooms']


@pytest.mark.parametrize('status', ['tentative', 'cancelled', 'released', 'expired'])
def test_non_definite_group_never_deducts_inventory(status):
    group = {'id': 'g', 'room_type': 'Standard', 'total_rooms': 10, 'status': status,
        'check_in': '2026-10-24', 'check_out': '2026-10-25'}
    row = report(groups=[group])['rows'][0]
    assert row['block_rooms'] == 0 and row['remaining_rooms'] == 60
    assert row['tentative_block_rooms'] == (10 if status == 'tentative' else 0)


def test_night_and_departure_boundaries_day_use_and_comp():
    departing = booking('depart', room_id='r0', status='checked_out', check_in='2026-10-23', checked_out_at='2026-10-24T09:00:00Z', check_out='2026-10-24')
    day_use = booking('day', room_id='r1', booking_type='day_use', check_out='2026-10-24', status='checked_in')
    comp = booking('comp', room_id='r2', is_complimentary=True)
    row = report([departing, day_use, comp])['rows'][0]
    assert row['booked_rooms'] == row['comp_rooms'] == row['day_use_rooms'] == row['arrivals'] == row['departures'] == 1
    assert row['room_revenue_by_currency'] == {'TRY': 0}


def test_historical_actuals_and_early_checkout_do_not_become_future_commitments():
    data = report([booking(room_id='r0', status='checked_out', check_out='2026-10-27', checked_out_at='2026-10-25T10:00:00Z')],
        first=DAY, last=date(2026, 10, 26), business_day=date(2026, 10, 25))
    assert [r['booked_rooms'] for r in data['rows']] == [1, 0, 0]
    assert [r['basis'] for r in data['rows']] == ['actual', 'on_the_books', 'on_the_books']


@pytest.mark.parametrize('status', ['cancelled', 'no_show', 'pending', 'tentative'])
def test_non_deducting_booking_statuses_excluded(status):
    row = report([booking(status=status)])['rows'][0]
    assert row['booked_rooms'] == row['arrivals'] == 0


def test_daily_rate_zero_currency_and_cent_preserving_fallback():
    data = report([booking('try', check_out='2026-10-27'), booking('eur', currency='EUR')],
        rates=[{'booking_id': 'eur', 'date': '2026-10-24', 'rate': 0}], last=date(2026, 10, 26))
    assert [r['room_revenue_by_currency']['TRY'] for r in data['rows']] == [33.33, 33.34, 33.33]
    assert data['rows'][0]['room_revenue_by_currency']['EUR'] == 0
    assert data['totals']['room_revenue_by_currency'] == {'EUR': 0, 'TRY': 100}
    assert data['totals']['adr_by_currency'] == {'EUR': 0, 'TRY': 33.33}


@pytest.mark.parametrize('rates', [[{'rate': 50}, {'rate': 60}], [{'rate': 'NaN'}], [{'rate': 50, 'currency': 'EUR'}]])
def test_incomplete_prices_never_look_like_zero_or_partial_total(rates):
    data = report([booking()], rates=[{'booking_id': 'b', 'date': '2026-10-24', **rate} for rate in rates])
    assert data['rows'][0]['room_revenue_by_currency'] is None
    assert data['totals']['room_revenue_by_currency'] is None
    assert data['totals']['adr_by_currency'] is None
    assert data['warnings']


def test_daily_complimentary_flag_and_missing_child_ages_are_explicit():
    data = report([booking(children=2, children_ages=[1])], rates=[{'booking_id': 'b', 'date': '2026-10-24', 'rate': 100, 'is_complimentary': True}])
    row = data['rows'][0]
    assert row['comp_rooms'] == 1 and row['room_revenue_by_currency'] == {'TRY': 0}
    assert row['children'] == 1 and row['infants'] is None and row['unknown_child_ages'] == 1
    malformed = report([booking(children=2, children_ages=123)])['rows'][0]
    assert malformed['children'] == 2 and malformed['infants'] is None and malformed['unknown_child_ages'] == 2


def test_overbooking_and_zero_capacity_not_hidden_or_clamped():
    data = report([booking(f'b{i}') for i in range(61)])
    assert data['rows'][0]['occupancy_pct'] == 101.67 and data['rows'][0]['remaining_rooms'] == -1
    empty = frontdesk_forecast([], [], [], [], [], DAY, DAY, DAY)
    assert empty['rows'][0]['occupancy_pct'] is None and empty['totals']['occupancy_pct'] is None


def test_unknown_inventory_cannot_show_a_false_exact_availability():
    group = {'id': 'g', 'total_rooms': 5, 'status': 'definite', 'check_in': '2026-10-24', 'check_out': '2026-10-25'}
    data = report([booking()], groups=[group], room_type='Standard')
    assert data['rows'][0]['booked_rooms'] == 1
    assert data['rows'][0]['remaining_rooms'] is None and data['rows'][0]['occupancy_pct'] is None
    assert data['totals']['occupied_room_nights'] is None
    bad_room = report([booking(room_id='missing')])
    assert bad_room['rows'][0]['inventory_incomplete'] is True
    assert bad_room['rows'][0]['remaining_rooms'] is None
    assert bad_room['rows'][0]['room_revenue_by_currency'] is None
    assert bad_room['totals']['room_revenue_by_currency'] is None


@pytest.mark.parametrize('view', ['availability', 'movements', 'occupancy', 'revenue', 'stay_types'])
def test_excel_is_same_data_numeric_not_formulas_and_matches_selected_view(view):
    data = report([booking()], room_type='Standard')
    data['room_type'] = '=1+1'
    book = load_workbook(forecast_workbook(data, view))
    assert book.active.title == 'Forecast'
    assert book.active.cell(7, 1).value.date() == DAY
    for sheet in book:
        assert all(cell.data_type != 'f' for row in sheet for cell in row)
    assert len(book['Müsaitlik Grafiği']._charts) == 1
    assert book['Müsaitlik Grafiği'].cell(2, 2).value == data['rows'][0]['booked_rooms']
    if view == 'availability':
        columns = {cell.value: cell.column for cell in book.active[6]}
        assert book.active.cell(7, columns['Doluluk % (fiziksel)']).value == data['rows'][0]['occupancy_pct'] / 100
        sellable = load_workbook(forecast_workbook(data, view, 'sellable'))
        assert any(cell.value == 'Doluluk % (satılabilir)' for cell in sellable.active[6])


@pytest.fixture
def api_dataset(monkeypatch):
    from routers.reports_pkg import frontdesk_forecast as api

    rows = {'rooms': ROOMS, 'bookings': [booking()], 'group_blocks': [], 'daily_rates': [], 'room_blocks': [], 'block_reservations': []}
    database = SimpleNamespace(**{key: SimpleNamespace(find=Mock(return_value=SimpleNamespace(to_list=AsyncMock(return_value=value)))) for key, value in rows.items()})
    monkeypatch.setattr(api, 'db', database)
    monkeypatch.setattr(api, 'ensure_business_date_initialized', AsyncMock(return_value={'business_date': DAY.isoformat()}))
    monkeypatch.setattr(api, 'get_tenant_currency', AsyncMock(return_value=('TRY', None)))
    return api, database


def test_excel_verified_empty_revenue_is_zero_but_missing_price_is_blank():
    data = report([booking()], last=date(2026, 10, 25))
    book = load_workbook(forecast_workbook(data, 'revenue'))
    columns = {cell.value: cell.column for cell in book.active[6]}
    assert book.active.cell(8, columns['Rezervasyon oda bedeli (TRY)']).value == 0
    data = report([booking(total_amount=None)])
    book = load_workbook(forecast_workbook(data, 'revenue'))
    assert book.active.cell(7, columns['Rezervasyon oda bedeli (TRY)']).value is None


async def read(api, **kwargs):
    return await api.get_frontdesk_forecast('2026-10-24', '2026-10-24', current_user=SimpleNamespace(tenant_id='A'), **kwargs)


@pytest.mark.asyncio
async def test_api_tenant_scope_no_pii_uncapped_and_snapshot_bound_export(api_dataset):
    api, database = api_dataset
    data = await read(api)
    for collection in vars(database).values():
        query, projection = collection.find.call_args.args
        assert query['tenant_id'] == 'A'
        assert not any(key in projection for key in ['guest_name', 'guest_id', 'email', 'phone', 'contact_name', 'contact_email'])
        collection.find.return_value.to_list.assert_awaited_once_with(None)
    response = await read(api, output='xlsx', snapshot=data['snapshot'])
    book = load_workbook(BytesIO(b''.join([chunk async for chunk in response.body_iterator])))
    assert book.active.cell(7, 4).value == data['rows'][0]['total_occupied_rooms']
    database.bookings.find.return_value.to_list.return_value = []
    with pytest.raises(HTTPException) as exc:
        await read(api, output='xlsx', snapshot=data['snapshot'])
    assert exc.value.status_code == 409


@pytest.mark.asyncio
async def test_invalid_scope_or_date_rejected(api_dataset):
    api, _ = api_dataset
    with pytest.raises(HTTPException) as exc:
        await read(api, room_type='missing')
    assert exc.value.status_code == 422
    with pytest.raises(HTTPException):
        await api.get_frontdesk_forecast('2026-02-30', '2026-10-24', current_user=SimpleNamespace(tenant_id='A'))


@pytest.mark.asyncio
async def test_legacy_allotment_and_booking_dates_share_nightly_pickup_without_counter_guessing(api_dataset):
    api, database = api_dataset
    database.block_reservations.find.return_value.to_list.return_value = [{'id': 'allotment', 'room_type': 'Standard',
        'start_date': '2026-10-24', 'end_date': '2026-10-25', 'total_rooms': 8, 'status': 'active', 'rooms_available': 0}]
    legacy = booking(block_id='allotment', check_in_date='2026-10-24', check_out_date='2026-10-25')
    del legacy['check_in'], legacy['check_out']
    database.bookings.find.return_value.to_list.return_value = [legacy]
    data = await read(api)
    row = data['rows'][0]
    assert (row['booked_rooms'], row['block_rooms'], row['group_rooms'], row['total_occupied_rooms']) == (1, 7, 1, 8)
    query = database.bookings.find.call_args.args[0]
    assert 'check_in_date' in str(query) and 'check_out_date' in str(query)
    database.block_reservations.find.return_value.to_list.return_value[0]['status'] = 'released'
    assert (await read(api))['rows'][0]['block_rooms'] == 0
