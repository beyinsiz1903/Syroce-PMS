from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from modules.pms_core.occupancy_calendar import occupancy_calendar
from modules.pms_core.occupancy_calendar_export import occupancy_summary, occupancy_workbook


def dataset():
    rooms = [{'id': str(i), 'room_type': 'Standard'} for i in range(10)]
    bookings = [
        {'id': 'stay', 'room_id': '0', 'status': 'checked_in', 'check_in': '2026-11-09', 'check_out': '2026-11-12'},
        {'id': 'arrival', 'room_id': '1', 'status': 'confirmed', 'check_in': '2026-11-10', 'check_out': '2026-11-12'},
    ]
    return occupancy_calendar(rooms, bookings, [], date(2026, 11, 9), date(2026, 11, 11), date(2026, 11, 10))


def test_summary_selects_one_metric_per_date_without_double_counting_business_day():
    data = dataset()
    rows, total = occupancy_summary(data)
    assert [row['basis'] for row in rows] == ['Gerçekleşen', 'Beklenen', 'Beklenen']
    assert [row['occupied'] for row in rows] == [1, 2, 2]
    assert [row['vacant'] for row in rows] == [9, 8, 8]
    assert total['occupancy'] == 5 / 30
    assert data['rows'][1]['actual_rooms'] == 1  # arrival is not a realised room


def test_weighted_summary_uses_room_nights_not_the_mean_of_daily_percentages():
    data = dataset()
    data['rows'][1].update(available_rooms=20, expected_rooms=14, expected_occupancy=70)
    _, total = occupancy_summary(data)
    assert total['occupancy'] == 17 / 40


def test_export_is_readable_typed_and_keeps_all_reconciliation_fields():
    data = dataset()
    workbook = load_workbook(occupancy_workbook(data))
    assert workbook.sheetnames == ['Doluluk Özeti', 'Detaylar']
    sheet = workbook.active
    assert sheet.max_column == 6
    assert [cell.value for cell in sheet[6]] == ['Tarih', 'Doluluk %', 'Dolu oda', 'Boş oda', 'Satılabilir oda', 'Veri türü']
    assert isinstance(sheet['A7'].value, datetime)
    assert sheet['A7'].number_format == 'dd.mm.yyyy'
    assert sheet['B8'].value == .2
    assert sheet['B8'].number_format == '0.00%'
    assert sheet['B4'].value == pytest.approx(5 / 30)
    assert sheet['D4'].value == 5 and sheet['F4'].value == 30
    assert sheet.freeze_panes == 'B7'
    assert sheet.auto_filter.ref == 'A6:F9'
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.sheet_view.showGridLines is False
    detail = workbook['Detaylar']
    assert detail.max_column == 15
    for index, row in enumerate(data['rows'], 7):
        assert detail.cell(index, 5).value == row['actual_rooms']
        assert detail.cell(index, 9).value == row['expected_rooms']
        assert detail.cell(index, 13).value == row['assignment_conflicts']
    assert detail['C11'].value == data['totals']['actual']['room_nights']
    assert detail['C12'].value == data['totals']['expected']['room_nights']


def test_zero_capacity_and_overbooking_are_not_disguised_as_valid_zeros():
    data = dataset()
    data['rows'][1].update(available_rooms=0, expected_rooms=2, expected_occupancy=None, overbooked_rooms=2)
    rows, _ = occupancy_summary(data)
    assert rows[1]['occupancy'] is None and rows[1]['vacant'] == 0
    data['rows'][2].update(available_rooms=1, expected_rooms=2, expected_occupancy=200, overbooked_rooms=1)
    sheet = load_workbook(occupancy_workbook(data)).active
    assert sheet['B8'].value is None and sheet['B9'].value == 2
    assert 'Uyarı: envanter' in sheet.cell(sheet.max_row, 1).value


def test_formula_looking_room_types_remain_text():
    data = dataset()
    data['room_type'] = '=HYPERLINK("https://example.invalid", "unsafe")'
    workbook = load_workbook(occupancy_workbook(data))
    for sheet in workbook:
        assert not any(cell.data_type == 'f' for row in sheet for cell in row)
    assert data['room_type'] in workbook.active['A2'].value
