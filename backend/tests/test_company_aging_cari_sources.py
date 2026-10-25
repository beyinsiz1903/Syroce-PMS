"""Presentation regression: a transferred folio is settled, its company debt is not."""
from copy import deepcopy
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from core import tenant_currency
from modules.pms_core.cari_receivable_reporting import account_positions, load_account_positions
from routers.departments import reports

pytest_plugins = ("tests.test_report_ledger_mongo_acceptance",)
DAY = date(2026, 10, 24)
USER = SimpleNamespace(tenant_id="A", role="admin", id="operator")


def account(**extra):
    return {"id": "a", "name": "Company", "currency": "TRY", "balance": 100, **extra}


def charge(**extra):
    return {"id": "c", "cari_account_id": "a", "transaction_type": "charge", "amount": 100,
        "currency": "TRY", "business_date": "2026-10-01", "due_date": "2026-10-31", **extra}


def positions(accounts, transactions, source="cari_accounts"):
    return account_positions(accounts, transactions, source=source, fallback_currency="TRY", today=DAY)


def test_account_debt_survives_its_folio_credit_and_uses_due_day():
    rows, failed, issues = positions([account()], [charge()])
    assert not failed and not issues
    assert sum(r["balance"] for r in rows) == 100
    assert rows[0]["age_days"] == 0 and not rows[0]["age_unknown"]


@pytest.mark.parametrize("kind", ["payment", "adjustment", "transfer_out"])
def test_unallocated_credit_reduces_balance_without_guessing_paid_invoice_age(kind):
    rows, failed, _ = positions([account(balance=70)], [charge(), charge(id="p", transaction_type=kind, amount=30)])
    assert not failed and rows[0]["balance"] == 70
    assert rows[0]["age_unknown"] and rows[0]["age_days"] is None


def test_explicit_credit_ages_only_the_remaining_named_debt():
    rows, failed, _ = positions([account(balance=120)], [charge(due_date="2026-10-01"), charge(id="later", amount=50),
        charge(id="credit", transaction_type="payment", amount=30, allocations=[{"transaction_id": "c", "amount": 30}])])
    assert not failed
    assert {(r["balance"], r["age_days"]) for r in rows} == {(70, 23), (50, 0)}


@pytest.mark.parametrize("parts", [[{"transaction_id": "missing", "amount": 30}], [{"transaction_id": "c", "amount": 31}],
    [{"transaction_id": "c", "amount": 20}, {"transaction_id": "c", "amount": 10}], [{"transaction_id": "c", "amount": "NaN"}],
    [{"transaction_id": "c", "amount": 30.001}], [{"transaction_id": "c", "amount": -30}], [None], []])
def test_invalid_credit_allocations_preserve_balance_but_never_guess_an_age(parts):
    rows, failed, _ = positions([account(balance=70)], [charge(), charge(id="credit", transaction_type="payment", amount=30, allocations=parts)])
    assert not failed and rows[0]["balance"] == 70 and rows[0]["age_unknown"]


def test_multiple_due_dates_keep_separate_amounts_not_oldest_age_for_all():
    rows, failed, _ = positions([account(balance=150)], [charge(due_date="2026-10-01"), charge(id="later", amount=50)])
    assert not failed
    assert {(r["balance"], r["age_days"]) for r in rows} == {(100, 23), (50, 0)}


def test_other_currency_credit_never_cancels_native_debt_or_creates_success_scalar():
    rows, failed, issues = positions([account()], [charge(), charge(id="p", transaction_type="payment", amount=100, currency="EUR")])
    assert failed == ["cari_accounts:a"]
    assert rows[0]["currency"] == "TRY" and rows[0]["balance"] == 100
    assert issues[0]["reason"] == "account_projection_differs_from_sources"


@pytest.mark.parametrize("status", ["pending", "failed", "voided", "cancelled"])
def test_uncommitted_or_cancelled_sources_are_not_debt(status):
    rows, failed, _ = positions([account(balance=0)], [charge(status=status)])
    assert rows == [] and not failed


@pytest.mark.parametrize("flag", ["voided", "deleted"])
def test_invalid_flag_excludes_transaction(flag):
    assert positions([account(balance=0)], [charge(**{flag: True})])[:2] == ([], [])


@pytest.mark.parametrize("amount", ["NaN", "Infinity", "bad", -1, 0])
def test_invalid_amount_is_incomplete_not_success_zero(amount):
    _, failed, issues = positions([account()], [charge(amount=amount)])
    assert failed and issues


def test_missing_source_history_preserves_legacy_debt_but_cannot_be_certified():
    rows, failed, _ = positions([account(balance=4501)], [])
    assert rows[0]["balance"] == 4501 and rows[0]["age_unknown"]
    assert failed == ["cari_accounts:a"]


def test_opening_balance_is_explicit_undated_debt_not_guessed_revenue():
    rows, failed, _ = positions([account(opening_balance=100)], [])
    assert not failed and rows[0]["balance"] == 100 and rows[0]["age_unknown"]


def test_missing_account_and_duplicate_transaction_identity_are_visible():
    assert positions([], [charge()])[1] == ["cari_accounts:a"]
    assert positions([account()], [charge(), charge()])[1] == ["cari_accounts:a"]


def test_bson_legacy_account_identity_and_no_source_mutation():
    from bson import ObjectId

    oid = ObjectId()
    accounts = [account(_id=oid)]
    transactions = [charge(cari_account_id=str(oid))]
    before = deepcopy((accounts, transactions))
    assert positions(accounts, transactions)[0][0]["balance"] == 100
    assert (accounts, transactions) == before


def test_city_ledger_is_a_distinct_source_with_its_own_id_namespace():
    rows, failed, _ = positions([account(account_name="Agency", current_balance=100)],
        [{**charge(), "account_id": "a"}], source="city_ledger_accounts")
    assert not failed and rows[0]["account_key"] == "city_ledger_accounts:a"


@pytest.fixture(autouse=True)
def context(monkeypatch):
    monkeypatch.setattr(reports, "ensure_business_date_initialized", AsyncMock(return_value={"business_date": DAY.isoformat()}))
    monkeypatch.setattr(reports, "_enforce", lambda *_: None)
    monkeypatch.setattr(tenant_currency, "get_tenant_currency", AsyncMock(return_value=("TRY", "₺")))


async def test_transfer_company_folio_residual_plus_cari_debt_counted_exactly_once(ledger_db, monkeypatch):
    monkeypatch.setattr(reports, "db", ledger_db)
    await ledger_db.folios.insert_one({"tenant_id": "A", "id": "f", "folio_type": "company", "status": "open", "currency": "TRY"})
    await ledger_db.folio_charges.insert_one({"tenant_id": "A", "folio_id": "f", "total": 150, "currency": "TRY"})
    await ledger_db.payments.insert_one({"tenant_id": "A", "folio_id": "f", "amount": 100, "currency": "TRY", "method": "city_ledger", "payment_type": "city_ledger_transfer", "status": "paid"})
    await ledger_db.cari_accounts.insert_one({"tenant_id": "A", **account()})
    await ledger_db.cari_transactions.insert_one({"tenant_id": "A", **charge()})
    aging = await reports._compute_company_aging("A")
    finance = await reports.get_finance_snapshot.__wrapped__(USER)
    assert aging["complete"] and aging["total_ar_by_currency"] == {"TRY": 150}
    assert finance["pending_ar"]["total"] == 150
    assert sum(r["folio_count"] for r in aging["companies"]) == 1
    response = await reports.export_company_aging_excel(USER)
    wb = load_workbook(BytesIO(b"".join([part async for part in response.body_iterator])))
    assert sum(row[3] for row in list(wb.active.values)[2:] if row[0] != "TOTAL") == 150
    assert wb.active.cell(2, 10).value == "Age Unknown"


async def test_three_currencies_and_both_account_sources_remain_separate(ledger_db, monkeypatch):
    monkeypatch.setattr(reports, "db", ledger_db)
    for i, code in enumerate(("TRY", "EUR", "USD"), 1):
        await ledger_db.cari_accounts.insert_one({"tenant_id": "A", **account(id=code, currency=code, balance=i * 100)})
        await ledger_db.cari_transactions.insert_one({"tenant_id": "A", **charge(id=code, cari_account_id=code, currency=code, amount=i * 100)})
    await ledger_db.city_ledger_accounts.insert_one({"tenant_id": "A", **account(current_balance=25, currency="EUR")})
    await ledger_db.city_ledger_transactions.insert_one({"tenant_id": "A", **charge(account_id="a", amount=25, currency="EUR")})
    await ledger_db.cari_accounts.insert_one({"tenant_id": "B", **account(balance=999999)})
    await ledger_db.cari_transactions.insert_one({"tenant_id": "B", **charge(amount=999999)})
    result = await reports._compute_company_aging("A")
    assert result["complete"] and result["total_ar"] is None
    assert result["total_ar_by_currency"] == {"TRY": 100, "EUR": 225, "USD": 300}
    assert result["company_count"] == 4


async def test_legacy_projection_mismatch_is_not_exported_as_complete(ledger_db, monkeypatch):
    monkeypatch.setattr(reports, "db", ledger_db)
    await ledger_db.cari_accounts.insert_one({"tenant_id": "A", **account()})
    result = await reports._compute_company_aging("A")
    assert result["complete"] is False and result["total_ar"] is None
    assert result["total_ar_by_currency"] == {"TRY": 100}
    assert result["companies"][0]["unaged_balance"] == 100
    with pytest.raises(HTTPException) as error:
        await reports.export_company_aging_excel(USER)
    assert error.value.status_code == 503


async def test_cross_collection_identity_collision_is_not_certified_as_two_accounts(ledger_db, monkeypatch):
    monkeypatch.setattr(reports, "db", ledger_db)
    await ledger_db.cari_accounts.insert_one({"tenant_id": "A", **account()})
    await ledger_db.cari_transactions.insert_one({"tenant_id": "A", **charge()})
    await ledger_db.city_ledger_accounts.insert_one({"tenant_id": "A", **account(current_balance=100)})
    await ledger_db.city_ledger_transactions.insert_one({"tenant_id": "A", **charge(account_id="a")})
    result = await reports._compute_company_aging("A")
    assert result["complete"] is False and result["total_ar"] is None
    assert result["failed_account_ids"] == ["cari_accounts:a", "city_ledger_accounts:a"]
    assert any(issue["reason"] == "ambiguous_cross_collection_account_identity" for issue in result["source_issues"])


async def test_account_sources_are_uncapped_and_tenant_scoped(monkeypatch):
    reads = []

    class Collection:
        def find(self, query):
            assert query == {"tenant_id": "A"}
            reads.append(query)
            return self

        async def to_list(self, count):
            assert count is None
            return []

    database = SimpleNamespace(**{key: Collection() for key in ("cari_accounts", "cari_transactions", "city_ledger_accounts", "city_ledger_transactions")})
    assert await load_account_positions(database, "A", fallback_currency="TRY", today=DAY) == ([], [], [])
    assert len(reads) == 4
