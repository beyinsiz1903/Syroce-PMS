"""Quality management: standards, audits, findings and CAPA follow-up."""

from __future__ import annotations

import csv
import io
import os
import re
import uuid
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Literal

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field, field_validator

from core.database import db
from core.helpers import create_audit_log
from core.security import get_current_user
from models.schemas import User
from modules.pms_core.role_permission_service import require_op

router = APIRouter(prefix="/api/quality", tags=["quality-management"])

KINDS = ("standard", "audit", "finding", "capa")
STATUSES = ("draft", "open", "in_progress", "pending_verification", "closed", "cancelled")
SEVERITIES = ("low", "medium", "high", "critical")
RECURRENCES = ("none", "weekly", "monthly", "quarterly", "yearly")
TRANSITIONS = {
    "draft": {"open", "cancelled"},
    "open": {"in_progress", "cancelled"},
    "in_progress": {"pending_verification", "cancelled"},
    "pending_verification": {"in_progress", "closed"},
    "closed": {"open"},
    "cancelled": {"open"},
}


def _now() -> str:
    return datetime.now(UTC).isoformat()


def _query(tenant_id: str, **extra: object) -> dict:
    return {"tenant_id": tenant_id, "_kind": {"$in": [f"quality_{kind}" for kind in KINDS]}, "deleted_at": {"$exists": False}, **extra}


class ChecklistItem(BaseModel):
    label: str = Field(min_length=2, max_length=240)
    required: bool = True
    result: Literal["pending", "pass", "fail", "not_applicable"] = "pending"
    note: str | None = Field(default=None, max_length=1000)


class QualityRecordInput(BaseModel):
    kind: Literal["standard", "audit", "finding", "capa"]
    title: str = Field(min_length=3, max_length=240)
    description: str | None = Field(default=None, max_length=4000)
    department: str = Field(min_length=2, max_length=100)
    severity: Literal["low", "medium", "high", "critical"] = "medium"
    owner_id: str | None = Field(default=None, max_length=120)
    due_at: str | None = None
    source: str | None = Field(default=None, max_length=120)
    source_record_id: str | None = Field(default=None, max_length=120)
    standard_id: str | None = Field(default=None, max_length=120)
    audit_id: str | None = Field(default=None, max_length=120)
    finding_id: str | None = Field(default=None, max_length=120)
    root_cause: str | None = Field(default=None, max_length=4000)
    corrective_action: str | None = Field(default=None, max_length=4000)
    preventive_action: str | None = Field(default=None, max_length=4000)
    verification_note: str | None = Field(default=None, max_length=4000)
    checklist: list[ChecklistItem] = Field(default_factory=list, max_length=100)
    evidence_urls: list[str] = Field(default_factory=list, max_length=20)
    recurrence: Literal["none", "weekly", "monthly", "quarterly", "yearly"] = "none"
    recurrence_until: str | None = None

    @field_validator("due_at")
    @classmethod
    def validate_due_at(cls, value: str | None) -> str | None:
        if not value:
            return None
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError("Geçersiz hedef tarihi") from exc
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=UTC)
        return parsed.astimezone(UTC).isoformat()

    @field_validator("recurrence_until")
    @classmethod
    def validate_recurrence_until(cls, value: str | None) -> str | None:
        if not value:
            return None
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError("Geçersiz tekrar bitiş tarihi") from exc
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=UTC)
        return parsed.astimezone(UTC).isoformat()


class TransitionInput(BaseModel):
    status: Literal["draft", "open", "in_progress", "pending_verification", "closed", "cancelled"]
    note: str | None = Field(default=None, max_length=2000)


class QualityDocumentInput(BaseModel):
    title: str = Field(min_length=3, max_length=240)
    document_type: Literal["procedure", "policy", "certificate", "external_audit", "other"]
    document_no: str | None = Field(default=None, max_length=100)
    revision: str | None = Field(default=None, max_length=50)
    issued_at: str | None = None
    expires_at: str | None = None
    owner_id: str | None = Field(default=None, max_length=120)
    notes: str | None = Field(default=None, max_length=2000)


def _advance_due(value: str, recurrence: str) -> str:
    current = datetime.fromisoformat(value.replace("Z", "+00:00"))
    days = {"weekly": 7, "monthly": 30, "quarterly": 90, "yearly": 365}[recurrence]
    return (current + timedelta(days=days)).astimezone(UTC).isoformat()


async def _materialize_recurring_audits(tenant_id: str, actor_id: str) -> int:
    """Create each due audit occurrence once and advance the template atomically."""
    now = _now()
    templates = await db.tasks.find(
        _query(tenant_id, _kind="quality_audit", recurrence={"$ne": "none"}, next_recurrence_at={"$lte": now}),
        {"_id": 0},
    ).limit(100).to_list(100)
    created = 0
    for template in templates:
        due = template.get("next_recurrence_at")
        recurrence = template.get("recurrence")
        if not due or recurrence not in RECURRENCES[1:]:
            continue
        if template.get("recurrence_until") and due > template["recurrence_until"]:
            await db.tasks.update_one(_query(tenant_id, id=template["id"]), {"$set": {"recurrence": "none", "updated_at": now}})
            continue
        occurrence_id = str(uuid.uuid4())
        occurrence = {
            **{k: v for k, v in template.items() if k not in {"id", "status", "history", "created_at", "updated_at", "next_recurrence_at"}},
            "id": occurrence_id,
            "title": f"{template['title']} · {due[:10]}",
            "status": "open",
            "due_at": due,
            "recurrence": "none",
            "template_id": template["id"],
            "created_at": now,
            "updated_at": now,
            "created_by": actor_id,
            "history": [{"at": now, "by": actor_id, "action": "generated_from_schedule", "template_id": template["id"]}],
        }
        next_due = _advance_due(due, recurrence)
        result = await db.tasks.update_one(
            _query(tenant_id, id=template["id"], next_recurrence_at=due),
            {"$set": {"next_recurrence_at": next_due, "updated_at": now}},
        )
        if result.modified_count:
            await db.tasks.insert_one(occurrence)
            created += 1
    return created


async def _sync_quality_alerts(tenant_id: str) -> None:
    now = _now()
    alert_query = _query(tenant_id, status={"$nin": ["closed", "cancelled"]})
    alert_query["$or"] = [{"due_at": {"$lt": now}}, {"severity": "critical"}]
    rows = await db.tasks.find(
        alert_query,
        {"_id": 0, "id": 1, "title": 1, "owner_id": 1, "due_at": 1, "severity": 1},
    ).to_list(1000)
    for row in rows:
        reason = "overdue" if row.get("due_at") and row["due_at"] < now else "critical"
        await db.quality_notifications.update_one(
            {"tenant_id": tenant_id, "record_id": row["id"], "reason": reason, "acknowledged_at": {"$exists": False}},
            {"$setOnInsert": {"id": str(uuid.uuid4()), "tenant_id": tenant_id, "record_id": row["id"], "title": row["title"], "owner_id": row.get("owner_id"), "reason": reason, "created_at": now}},
            upsert=True,
        )


def validate_transition(current: str, target: str, record: dict) -> None:
    if target not in TRANSITIONS.get(current, set()):
        raise HTTPException(status_code=409, detail=f"'{current}' durumundan '{target}' durumuna geçilemez")
    if target == "closed":
        if record.get("_kind") == "quality_capa" and not record.get("verification_note"):
            raise HTTPException(status_code=422, detail="CAPA kapatılmadan önce doğrulama notu zorunludur")
        failed_required = [item for item in record.get("checklist", []) if item.get("required") and item.get("result") in {"pending", "fail"}]
        if failed_required:
            raise HTTPException(status_code=422, detail="Zorunlu kontrol maddeleri tamamlanmadan kayıt kapatılamaz")


@router.get("/dashboard")
async def quality_dashboard(
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_quality")),
):
    generated = await _materialize_recurring_audits(current_user.tenant_id, current_user.id)
    await _sync_quality_alerts(current_user.tenant_id)
    rows = await db.tasks.find(_query(current_user.tenant_id), {"_id": 0}).sort("created_at", -1).to_list(2000)
    now = _now()
    summary = {"total": len(rows), "open": 0, "overdue": 0, "critical": 0, "closed": 0, "by_kind": dict.fromkeys(KINDS, 0), "by_department": {}}
    for row in rows:
        status = row.get("status", "open")
        kind = str(row.get("_kind", "")).removeprefix("quality_")
        if kind in summary["by_kind"]:
            summary["by_kind"][kind] += 1
        department = row.get("department") or "Diğer"
        summary["by_department"][department] = summary["by_department"].get(department, 0) + 1
        if status == "closed":
            summary["closed"] += 1
        elif status != "cancelled":
            summary["open"] += 1
            if row.get("due_at") and row["due_at"] < now:
                summary["overdue"] += 1
        if row.get("severity") == "critical" and status not in {"closed", "cancelled"}:
            summary["critical"] += 1
    summary["closure_rate"] = round(summary["closed"] / summary["total"] * 100, 1) if summary["total"] else 0
    return {"summary": summary, "recent": rows[:10], "generated_audits": generated}


@router.get("/records")
async def list_quality_records(
    kind: Literal["standard", "audit", "finding", "capa"] | None = None,
    status: Literal["draft", "open", "in_progress", "pending_verification", "closed", "cancelled"] | None = None,
    department: str | None = None,
    limit: int = Query(200, ge=1, le=500),
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_quality")),
):
    query = _query(current_user.tenant_id)
    if kind:
        query["_kind"] = f"quality_{kind}"
    if status:
        query["status"] = status
    if department:
        query["department"] = department
    rows = await db.tasks.find(query, {"_id": 0}).sort("updated_at", -1).limit(limit).to_list(limit)
    return {"records": rows, "total": len(rows)}


@router.post("/records", status_code=201)
async def create_quality_record(
    payload: QualityRecordInput,
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("manage_quality")),
):
    now = _now()
    record = {
        **payload.model_dump(),
        "_kind": f"quality_{payload.kind}",
        "id": str(uuid.uuid4()),
        "tenant_id": current_user.tenant_id,
        "status": "draft" if payload.kind == "standard" else "open",
        "created_at": now,
        "updated_at": now,
        "created_by": current_user.id,
        "history": [{"at": now, "by": current_user.id, "action": "created"}],
    }
    if payload.kind == "audit" and payload.recurrence != "none" and payload.due_at:
        record["next_recurrence_at"] = payload.due_at
    record.pop("kind", None)
    await db.tasks.insert_one(record)
    await create_audit_log(current_user.tenant_id, current_user, "quality_record_created", "quality_record", record["id"], {"kind": payload.kind})
    return {"record": {k: v for k, v in record.items() if k != "_id"}}


@router.put("/records/{record_id}")
async def update_quality_record(
    record_id: str,
    payload: QualityRecordInput,
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("manage_quality")),
):
    existing = await db.tasks.find_one(_query(current_user.tenant_id, id=record_id), {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Kalite kaydı bulunamadı")
    if existing.get("_kind") != f"quality_{payload.kind}":
        raise HTTPException(status_code=409, detail="Kayıt türü değiştirilemez")
    update = payload.model_dump()
    update.pop("kind", None)
    update.update({"updated_at": _now(), "updated_by": current_user.id})
    update["next_recurrence_at"] = payload.due_at if payload.kind == "audit" and payload.recurrence != "none" and payload.due_at else None
    await db.tasks.update_one(_query(current_user.tenant_id, id=record_id), {"$set": update})
    await create_audit_log(current_user.tenant_id, current_user, "quality_record_updated", "quality_record", record_id)
    return {"success": True, "record_id": record_id}


@router.post("/records/{record_id}/transition")
async def transition_quality_record(
    record_id: str,
    payload: TransitionInput,
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("manage_quality")),
):
    record = await db.tasks.find_one(_query(current_user.tenant_id, id=record_id), {"_id": 0})
    if not record:
        raise HTTPException(status_code=404, detail="Kalite kaydı bulunamadı")
    current = record.get("status", "open")
    validate_transition(current, payload.status, record)
    now = _now()
    event = {"at": now, "by": current_user.id, "action": "status_changed", "from": current, "to": payload.status, "note": payload.note}
    await db.tasks.update_one(_query(current_user.tenant_id, id=record_id), {"$set": {"status": payload.status, "updated_at": now, "updated_by": current_user.id}, "$push": {"history": event}})
    await create_audit_log(current_user.tenant_id, current_user, "quality_status_changed", "quality_record", record_id, {"from": current, "to": payload.status})
    return {"success": True, "record_id": record_id, "status": payload.status}


@router.delete("/records/{record_id}")
async def archive_quality_record(
    record_id: str,
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("manage_quality")),
):
    now = _now()
    result = await db.tasks.update_one(_query(current_user.tenant_id, id=record_id), {"$set": {"deleted_at": now, "deleted_by": current_user.id}})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Kalite kaydı bulunamadı")
    await create_audit_log(current_user.tenant_id, current_user, "quality_record_archived", "quality_record", record_id)
    return {"success": True, "record_id": record_id}


@router.get("/feedback")
async def quality_feedback(
    limit: int = Query(50, ge=1, le=200),
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_quality")),
):
    rows = await db.feedback_entries.find({"tenant_id": current_user.tenant_id}, {"_id": 0}).sort("responded_at", -1).limit(limit).to_list(limit)
    ratings = [float(row["rating"]) for row in rows if row.get("rating") is not None]
    detractors = sum(1 for row in rows if row.get("nps_eligible") and int(row.get("nps_score", 0)) <= 6)
    return {
        "entries": rows,
        "summary": {
            "total": len(rows),
            "average_rating": round(sum(ratings) / len(ratings), 2) if ratings else None,
            "detractors": detractors,
            "unresolved": sum(1 for row in rows if row.get("resolution_status") not in {"resolved", "closed"}),
        },
    }


@router.get("/staff")
async def quality_staff(
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_quality")),
):
    """Minimal tenant-scoped assignee directory; never exposes HR/finance PII."""
    tenant_id = current_user.tenant_id
    staff = await db.staff_members.find(
        {"tenant_id": tenant_id, "active": {"$ne": False}},
        {"_id": 0, "id": 1, "name": 1, "department": 1, "position": 1},
    ).sort("name", 1).to_list(1000)
    users = await db.users.find(
        {"tenant_id": tenant_id, "is_active": {"$ne": False}},
        {"_id": 0, "id": 1, "full_name": 1, "name": 1, "department": 1, "role": 1},
    ).sort("full_name", 1).to_list(500)
    seen = {row.get("id") for row in staff}
    staff.extend({"id": row.get("id"), "name": row.get("full_name") or row.get("name") or row.get("role"), "department": row.get("department"), "position": row.get("role")} for row in users if row.get("id") not in seen)
    return {"staff": staff}


@router.get("/notifications")
async def quality_notifications(
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_quality")),
):
    await _sync_quality_alerts(current_user.tenant_id)
    rows = await db.quality_notifications.find(
        {"tenant_id": current_user.tenant_id, "acknowledged_at": {"$exists": False}}, {"_id": 0}
    ).sort("created_at", -1).limit(200).to_list(200)
    return {"notifications": rows}


@router.post("/notifications/{notification_id}/acknowledge")
async def acknowledge_quality_notification(
    notification_id: str,
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("manage_quality")),
):
    result = await db.quality_notifications.update_one(
        {"tenant_id": current_user.tenant_id, "id": notification_id},
        {"$set": {"acknowledged_at": _now(), "acknowledged_by": current_user.id}},
    )
    if not result.matched_count:
        raise HTTPException(status_code=404, detail="Bildirim bulunamadı")
    return {"success": True}


@router.post("/records/{record_id}/evidence", status_code=201)
async def upload_quality_evidence(
    record_id: str,
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("manage_quality")),
):
    record = await db.tasks.find_one(_query(current_user.tenant_id, id=record_id), {"_id": 0, "evidence_urls": 1})
    if not record:
        raise HTTPException(status_code=404, detail="Kalite kaydı bulunamadı")
    if len(record.get("evidence_urls", [])) >= 20:
        raise HTTPException(status_code=409, detail="Bir kayda en fazla 20 kanıt eklenebilir")
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Dosya çok büyük (en fazla 10 MB)")
    is_pdf = content.startswith(b"%PDF-")
    ext = ".pdf"
    content_type = "application/pdf"
    if not is_pdf:
        from security.upload_validator import validate_image_bytes
        content_type, ext = validate_image_bytes(content, max_bytes=10 * 1024 * 1024, field_label="Kanıt")
    safe_id = re.sub(r"[^a-zA-Z0-9_-]", "", record_id)
    upload_root = Path(os.environ.get("UPLOAD_DIR", str(Path(__file__).resolve().parents[2] / "uploads"))).resolve()
    target_dir = upload_root / current_user.tenant_id / "quality" / safe_id
    target_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{uuid.uuid4().hex}{ext}"
    target = (target_dir / filename).resolve()
    if upload_root not in target.parents:
        raise HTTPException(status_code=400, detail="Geçersiz dosya yolu")
    target.write_bytes(content)
    url = f"/api/uploads/{current_user.tenant_id}/quality/{safe_id}/{filename}"
    evidence = {"id": str(uuid.uuid4()), "url": url, "name": (file.filename or f"kanıt{ext}")[:180], "content_type": content_type, "size": len(content), "uploaded_at": _now(), "uploaded_by": current_user.id}
    await db.tasks.update_one(_query(current_user.tenant_id, id=record_id), {"$push": {"evidence_urls": url, "evidence": evidence}, "$set": {"updated_at": _now()}})
    return {"evidence": evidence}


@router.get("/export.csv")
async def export_quality_csv(
    current_user: User = Depends(get_current_user),
    _perm=Depends(require_op("view_quality")),
):
    rows = await db.tasks.find(_query(current_user.tenant_id), {"_id": 0}).sort("created_at", -1).to_list(5000)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Tür", "Başlık", "Departman", "Sorumlu", "Önem", "Durum", "Hedef tarih", "Kapanış tarihi"])
    for row in rows:
        writer.writerow([str(row.get("_kind", "")).removeprefix("quality_"), row.get("title"), row.get("department"), row.get("owner_id"), row.get("severity"), row.get("status"), row.get("due_at"), row.get("closed_at")])
    await create_audit_log(current_user.tenant_id, current_user, "quality_records_exported", "quality_record", "all", {"count": len(rows)})
    return StreamingResponse(iter(["\ufeff" + output.getvalue()]), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": 'attachment; filename="quality-register.csv"'})


async def _quality_export_rows(tenant_id: str) -> list[dict]:
    return await db.tasks.find(_query(tenant_id), {"_id": 0}).sort("created_at", -1).to_list(5000)


@router.get("/export.xlsx")
async def export_quality_xlsx(current_user: User = Depends(get_current_user), _perm=Depends(require_op("view_quality"))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    rows = await _quality_export_rows(current_user.tenant_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kalite Kayıtları"
    headings = ["Tür", "Başlık", "Departman", "Sorumlu", "Önem", "Durum", "Hedef tarih", "Kanıt sayısı"]
    sheet.append(headings)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for row in rows:
        sheet.append([str(row.get("_kind", "")).removeprefix("quality_"), row.get("title"), row.get("department"), row.get("owner_id"), row.get("severity"), row.get("status"), row.get("due_at"), len(row.get("evidence_urls", []))])
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(48, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="quality-register.xlsx"'})


@router.get("/export.pdf")
async def export_quality_pdf(current_user: User = Depends(get_current_user), _perm=Depends(require_op("view_quality"))):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    rows = await _quality_export_rows(current_user.tenant_id)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title="Quality Register")
    data = [["Type", "Title", "Department", "Severity", "Status", "Due"]]
    data.extend([[str(row.get("_kind", "")).removeprefix("quality_"), str(row.get("title", ""))[:60], row.get("department", ""), row.get("severity", ""), row.get("status", ""), str(row.get("due_at") or "")[:10]] for row in rows])
    table = Table(data, repeatRows=1, colWidths=[70, 260, 110, 70, 100, 75])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#cbd5e1")), ("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    doc.build([table])
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": 'attachment; filename="quality-register.pdf"'})


@router.get("/documents")
async def list_quality_documents(current_user: User = Depends(get_current_user), _perm=Depends(require_op("view_quality"))):
    rows = await db.quality_documents.find({"tenant_id": current_user.tenant_id, "deleted_at": {"$exists": False}}, {"_id": 0}).sort("updated_at", -1).to_list(1000)
    return {"documents": rows}


@router.post("/documents", status_code=201)
async def create_quality_document(payload: QualityDocumentInput, current_user: User = Depends(get_current_user), _perm=Depends(require_op("manage_quality"))):
    now = _now()
    document = {**payload.model_dump(), "id": str(uuid.uuid4()), "tenant_id": current_user.tenant_id, "status": "active", "created_at": now, "updated_at": now, "created_by": current_user.id, "revision_history": [{"revision": payload.revision or "1", "at": now, "by": current_user.id}]}
    await db.quality_documents.insert_one(document)
    await create_audit_log(current_user.tenant_id, current_user, "quality_document_created", "quality_document", document["id"])
    return {"document": {k: v for k, v in document.items() if k != "_id"}}
